"""Build v3.5 from v3.4: theory-based capital market assumptions in the Forecast tab.

Usage: python build_v35.py <v3.4.xlsx> <out.xlsx>

Methods (sources are written next to every input in Forecast section 5):
  Equities          Grinold-Kroner (2002), as in the CFA Institute curriculum: dividend yield - net dilution
                    + real earnings growth + inflation + repricing; foreign assets in AUD via relative PPP.
                    Averaged with a risk-premium build-up (expected cash + Dimson-Marsh-Staunton premium).
  Listed property   Grinold-Kroner on REIT yields; build-up with a data-estimated beta.
  Direct property   CFA cap-rate model: net rental yield + NOI growth - change in cap rate.
  Cash              Expectations hypothesis: 5-year yield less a term premium, averaged with the current rate.
  Australian bonds  Starting yield to maturity less expected credit losses (Leibowitz; Lozada 2015).
  World bonds (hdg) Covered interest parity: foreign yields + AUD/foreign cash-rate gap - basis.
  Commodities       Collateral + long-run excess return (Erb & Harvey 2006; Levine et al. 2018) + PPP drift.
  Hedge funds       Cash + equity beta x premium, zero net alpha (Fung & Hsieh 2004; Dichev & Yu 2011).
  Private equity    Cash + levered beta x premium + net premium (Harris, Jenkinson & Kaplan 2014).
All views are compound (geometric) returns; they are converted to arithmetic (+ sigma^2/2) before the
Black-Litterman-style blend because the optimiser and the equilibrium prior are arithmetic.
"""
import sys

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from style import (BLUE, BROWN, GOLD, GREEN, GREY, NAVY, OLIVE, ORANGE, PURPLE, RED, TEAL, box,
                   status_cf)

ASSETS = ["Australian Equities", "World Equities", "Emerging Markets", "Listed Property",
          "Commodities", "Direct Property", "Hedge Funds", "Private Equity",
          "Australian Fixed Income", "World Fixed Income", "Cash"]
PCT1, PCT2 = "0.0%", "0.00%"

# v3.4 net forecasts, kept for the comparison table
V34_NET = [0.08268, 0.08218, 0.07880, 0.07320, 0.05030, 0.05440, 0.03320, 0.05870, 0.04790,
           0.04770, 0.04600]


def merge_text(ws, r, c0, c1, height=None):
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
    ws.cell(r, c0).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if height:
        ws.row_dimensions[r].height = height


# ---------------------------------------------------------------------------------------------
def quarterly_helpers(wb):
    """Excess returns over cash for the factor regressions (adjusted quarterly returns, rows 6-85)."""
    q = wb["Quarterly returns"]
    names = [("World equities excess", "O"), ("Listed property excess", "Q"),
             ("Direct property excess", "S"), ("Private equity excess", "U"),
             ("Hedge funds excess", "T")]
    rows = [[None] * 5] + [[f"={c}{r}-$X{r}" for _, c in names] for r in range(6, 86)]
    box(q, 3, 32, "Excess returns over cash (adjusted; for Forecast section 6c)", [n for n, _ in names],
        rows, GREY, fmts=[PCT2] * 5, header_height=28)
    for i in range(5):
        q.column_dimensions[L(32 + i)].width = 11
    q.column_dimensions["AE"].width = 3
    return {n: L(32 + i) for i, (n, _) in enumerate(names)}


def run(wb):
    f = wb["Forecast"]
    QC = quarterly_helpers(wb)

    # ---- 5. Market and economic inputs ----
    top = 116
    items = [
        ("cash", "AU cash rate (RBA target)", "=Inputs!$B$26", "RBA, 29-Sep-26 (Inputs B26)", False, PCT2),
        ("y5", "AU 5-year government bond yield", "=Inputs!$B$23", "RBA F2, 22-Sep-26 (Inputs B23)", False, PCT2),
        ("tp", "AU 5-year term premium", 0.003, "Judgement; Hambur & Finlay (2018), RBA RDP 2018-02, find small positive term premia", True, PCT2),
        ("piA", "Australian inflation, 5-year average", 0.0275, "RBA Statement on Monetary Policy, Aug-2026: inflation back to 2.5% by early 2028", True, PCT2),
        ("piW", "Foreign inflation, 5-year average", 0.025, "Judgement: above 2% targets near term (Fed, ECB and BoJ all raised rates in Sep-2026)", True, PCT2),
        ("dms", "Equity premium over bills (geometric, long run)", 0.035, "Dimson, Marsh & Staunton, UBS Global Investment Returns Yearbook 2026: c.3.5% p.a.", True, PCT2),
        ("rev", "Share of the P/E gap closed over 5 years", 0.5, "Judgement; Campbell & Shiller (1998): valuations mean-revert slowly", True, "0%"),
        ("ae_dy", "Australian equities: forward dividend yield", "=$B$46", "Section 3 (market data, 31-Aug-26)", False, PCT2),
        ("ae_pe", "Australian equities: forward P/E", "=$B$43", "Section 3 (market data, 31-Aug-26)", False, "0.0"),
        ("ae_pelr", "Australian equities: long-run average forward P/E", 14.8, "IG, ASX 200 outlook 2026 (verify: J.P. Morgan Guide to the Markets Australia)", True, "0.0"),
        ("ae_g", "Australian equities: real earnings growth", 0.02, "Judgement: about potential GDP growth (RBA SMP Aug-2026)", True, PCT2),
        ("ae_dil", "Australian equities: net dilution", 0.01, "Bernstein & Arnott (2003): c.2% p.a.; judgement 1% (DRPs offset by buybacks)", True, PCT2),
        ("we_dy", "World equities: forward dividend yield", "=$C$46", "Section 3 (market data, 31-Aug-26)", False, PCT2),
        ("we_pe", "World equities: forward P/E", "=$C$43", "Section 3 (market data, 31-Aug-26)", False, "0.0"),
        ("we_pelr", "World equities: long-run average forward P/E", 15.5, "Judgement (verify: J.P. Morgan Guide to the Markets, MSCI World 20-year average)", True, "0.0"),
        ("we_g", "World equities: real earnings growth", 0.018, "Judgement: developed-market potential growth", True, PCT2),
        ("we_dil", "World equities: net dilution", 0.005, "Judgement: US buybacks offset issuance elsewhere; Bernstein & Arnott (2003)", True, PCT2),
        ("em_dy", "Emerging markets: dividend yield", 0.02, "iShares MSCI EM ETF trailing yield, May-2026 (verify: MSCI EM index factsheet)", True, PCT2),
        ("em_pe", "Emerging markets: forward P/E", 11.6, "Siblis Research, Jan-2026", True, "0.0"),
        ("em_pelr", "Emerging markets: long-run average forward P/E", 12.0, "Judgement (verify: J.P. Morgan Guide to the Markets)", True, "0.0"),
        ("em_g", "Emerging markets: real earnings growth", 0.035, "Judgement: EM potential growth", True, PCT2),
        ("em_dil", "Emerging markets: net dilution", 0.02, "Bernstein & Arnott (2003): c.2% p.a.", True, PCT2),
        ("reit_dy", "Listed property: REIT dividend yield", 0.0368, "Nareit, FTSE Nareit All Equity REITs, Aug-2026", True, PCT2),
        ("reit_dil", "Listed property: net issuance", 0.01, "Judgement: REITs fund growth with new equity", True, PCT2),
        ("reit_g", "Listed property: real rent growth", 0.005, "Judgement", True, PCT2),
        ("res_yld", "Direct property: gross rental yield", 0.0379, "Cotality, national dwellings, Aug-2026", True, PCT2),
        ("res_run", "Direct property: running costs (% of value)", 0.015, "Fox & Tulip (2014), RBA RDP 2014-06", True, PCT2),
        ("res_dep", "Direct property: depreciation (% of value)", 0.011, "Fox & Tulip (2014), RBA RDP 2014-06", True, PCT2),
        ("res_g", "Direct property: real rent growth", 0.005, "Judgement", True, PCT2),
        ("afi_ytm", "Australian fixed income: yield to maturity", 0.051, "Vanguard VAF fact sheet 2026 (Bloomberg AusBond Composite)", True, PCT2),
        ("afi_dur", "Australian fixed income: duration (years)", 4.9, "Vanguard VAF fact sheet 2026", True, "0.0"),
        ("afi_cl", "Australian fixed income: expected credit losses", 0.0005, "Judgement: AA+ average quality", True, PCT2),
        ("us_y", "US 10-year Treasury yield", 0.0508, "Weekly average, end Sep-2026 (FRED DGS10 / Advisor Perspectives)", True, PCT2),
        ("us_i", "US policy rate (midpoint)", 0.03875, "FOMC, 16-Sep-26: 3.75% to 4.00%", True, PCT2),
        ("eu_y", "Germany 10-year Bund yield", 0.036, "29-Sep-26 (Trading Economics)", True, PCT2),
        ("eu_i", "ECB deposit facility rate", 0.025, "ECB, 10-Sep-26", True, PCT2),
        ("jp_y", "Japan 10-year JGB yield", 0.0308, "28-Sep-26 (Trading Economics)", True, PCT2),
        ("jp_i", "Bank of Japan policy rate", 0.0125, "Bank of Japan, 18-Sep-26", True, PCT2),
        ("uk_y", "UK 10-year gilt yield", 0.0538, "25-Sep-26 (Trading Economics)", True, PCT2),
        ("uk_i", "Bank of England Bank Rate", 0.0375, "Bank of England, Sep-2026", True, PCT2),
        ("w_us", "World bonds: United States weight", 0.397, "Vanguard VIF fact sheet, country allocation", True, PCT1),
        ("w_jp", "World bonds: Japan weight", 0.155, "Vanguard VIF fact sheet", True, PCT1),
        ("w_uk", "World bonds: United Kingdom weight", 0.061, "Vanguard VIF fact sheet", True, PCT1),
        ("w_eu", "World bonds: euro area and other weight", None, "Remainder, priced at Bund rates (conservative)", False, PCT1),
        ("hdg_adj", "World bonds: duration and basis adjustment", -0.0025, "Judgement: 10-year yields overstate a 6.7-year duration index (-0.15%); basis and costs (-0.10%)", True, PCT2),
        ("com_x", "Commodities: excess return over collateral", 0.015, "Levine et al. (2018): 3.3% geometric for a diversified index; Erb & Harvey (2006): c.0% for the average commodity", True, PCT2),
        ("hf_beta", "Hedge funds: equity beta", 0.35, "Judgement from factor-model literature (Fung & Hsieh 2004); data estimate in section 6c", True, "0.00"),
        ("hf_alpha", "Hedge funds: net alpha", 0.0, "Dichev & Yu (2011): investors earn 3-7% less than fund returns", True, PCT2),
        ("pe_beta", "Private equity: equity beta", 1.2, "Korteweg (2019) survey: 0.7 to 3.2; Axelson, Sorensen & Stromberg (2014); judgement", True, "0.00"),
        ("pe_prem", "Private equity: net premium over levered public equity", 0.01, "Harris, Jenkinson & Kaplan (2014): >3% p.a. (1984-2008 vintages); judgement 1% for recent vintages", True, PCT2),
    ]
    K = {}
    rows = []
    for i, (key, label, val, srcs, is_in, nf) in enumerate(items):
        K[key] = f"$B${top + 2 + i}"
    for i, (key, label, val, srcs, is_in, nf) in enumerate(items):
        if key == "w_eu":
            val = f"=1-{K['w_us']}-{K['w_jp']}-{K['w_uk']}"
        rows.append([label, val, srcs] + [None] * 6)
    last5 = box(f, top, 1, "5. Market and economic inputs, Sep-2026 (blue = sourced or judgement input)",
                ["Input", "Value", "Source"] + [None] * 6, rows, BLUE,
                fmts=lambda ri, ci: items[ri][5] if ci == 1 else None,
                inputs={(i, 1) for i, it in enumerate(items) if it[4]}, notes_col=2)
    for r in range(top + 1, last5 + 1):
        merge_text(f, r, 3, 9)
        f.cell(r, 3).font = f.cell(r, 3).font.copy(size=9)

    # ---- 6a. Yield, carry and factor builds ----
    top6 = last5 + 3
    r0 = top6 + 2
    R = {a: r0 + i for i, a in enumerate(["Cash", "Australian Fixed Income", "World Fixed Income",
                                           "Commodities", "Direct Property", "Hedge Funds", "Private Equity"])}
    # world equities composite (compound) lives in section 7; defined after layout below
    top6b = top6 + 2 + len(R) + 3
    E0 = top6b + 2
    RE = {a: E0 + i for i, a in enumerate(["Australian Equities", "World Equities", "Emerging Markets",
                                           "Listed Property"])}
    top6c = E0 + len(RE) + 3
    C0 = top6c + 2
    RB = {a: C0 + i for i, a in enumerate(["Listed Property", "Direct Property", "Private Equity", "Hedge Funds"])}
    top7 = C0 + len(RB) + 3
    S0 = top7 + 2
    RS = {a: S0 + i for i, a in enumerate(ASSETS)}
    cash_exp = f"$D${RS['Cash']}"          # composite (compound) of the cash views
    we_comp = f"$D${RS['World Equities']}"

    wy = f"({K['w_us']}*{K['us_y']}+{K['w_jp']}*{K['jp_y']}+{K['w_uk']}*{K['uk_y']}+{K['w_eu']}*{K['eu_y']})"
    wi = f"({K['w_us']}*{K['us_i']}+{K['w_jp']}*{K['jp_i']}+{K['w_uk']}*{K['uk_i']}+{K['w_eu']}*{K['eu_i']})"
    b_dp = f"$B${RB['Direct Property']}"
    b_lp = f"$B${RB['Listed Property']}"
    rows = []
    spec = {
        "Cash": [f"={K['y5']}", None, f"=-{K['tp']}", None, "Expectations hypothesis", f"={K['cash']}"],
        "Australian Fixed Income": [f"={K['afi_ytm']}", None, f"=-{K['afi_cl']}", None, "Yield to maturity", None],
        "World Fixed Income": [f"={wy}", None, f"={K['hdg_adj']}", f"={K['cash']}-{wi}", "Covered interest parity", None],
        "Commodities": [f"={K['us_i']}", f"={K['com_x']}", None, f"={K['piA']}-{K['piW']}", "Collateral + excess + PPP", None],
        "Direct Property": [f"={K['res_yld']}-{K['res_run']}-{K['res_dep']}", f"={K['piA']}+{K['res_g']}", 0, None, "Cap-rate model; B: beta build-up", f"={cash_exp}+{b_dp}*{K['dms']}"],
        "Hedge Funds": [f"={cash_exp}", f"={K['hf_beta']}*({we_comp}-{cash_exp})", f"={K['hf_alpha']}", None, "Factor model (net of fees)", None],
        "Private Equity": [f"={cash_exp}", f"={K['pe_beta']}*({we_comp}-{cash_exp})", f"={K['pe_prem']}", None, "Levered equity (net of fees)", None],
    }
    for a, rr in R.items():
        s = spec[a]
        rows.append([a, s[0], s[1], s[2], s[3], f"=SUM(B{rr}:E{rr})", s[5], s[4]])
    last6 = box(f, top6, 1, "6a. Yield, carry and factor builds (compound, p.a.)",
                ["Asset class", "Starting yield or income", "Growth or carry", "Adjustments",
                 "Currency or hedge", "View A", "View B", "Method"], rows, TEAL,
                fmts=[None, PCT2, PCT2, PCT2, PCT2, PCT2, PCT2, None], header_height=40, notes_col=7)
    for rr in R.values():
        f.cell(rr, 8).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ---- 6b. Grinold-Kroner equities ----
    gk = {
        "Australian Equities": ("ae_dy", "ae_dil", "ae_g", "ae_pe", "ae_pelr", "1"),
        "World Equities": ("we_dy", "we_dil", "we_g", "we_pe", "we_pelr", "1"),
        "Emerging Markets": ("em_dy", "em_dil", "em_g", "em_pe", "em_pelr", "1"),
        "Listed Property": ("reit_dy", "reit_dil", "reit_g", None, None, b_lp),
    }
    rows = []
    for a, rr in RE.items():
        dy, dil, g, pe, pelr, beta = gk[a]
        rep = f"=({K[pelr]}/{K[pe]})^({K['rev']}/5)-1" if pe else 0
        rows.append([a, f"={K[dy]}", f"=-{K[dil]}", f"={K[g]}", f"={K['piA']}", rep,
                     f"=SUM(B{rr}:F{rr})", f"={beta}", f"={cash_exp}+H{rr}*{K['dms']}"])
    box(f, top6b, 1, "6b. Equities and listed property: Grinold-Kroner and risk-premium build-up "
        "(compound, p.a.; foreign returns in AUD via relative PPP)",
        ["Asset class", "Dividend yield", "Net dilution", "Real earnings growth",
         "Inflation (AUD)", "Repricing", "View A: Grinold-Kroner", "Beta to equity premium",
         "View B: cash + beta x premium"], rows, NAVY,
        fmts=[None, PCT2, PCT2, PCT2, PCT2, PCT2, PCT2, "0.00", PCT2], header_height=40)

    # ---- 6c. Regression betas ----
    qx = {"World": QC["World equities excess"], "Listed Property": QC["Listed property excess"],
          "Direct Property": QC["Direct property excess"], "Private Equity": QC["Private equity excess"],
          "Hedge Funds": QC["Hedge funds excess"]}
    rows = []
    lit = {"Listed Property": None, "Direct Property": None, "Private Equity": f"={K['pe_beta']}",
           "Hedge Funds": f"={K['hf_beta']}"}
    used = {"Listed Property": "Yes", "Direct Property": "Yes",
            "Private Equity": "No: AUD index distorts beta", "Hedge Funds": "No: AUD index distorts beta"}
    for a, rr in RB.items():
        y = f"'Quarterly returns'!${qx[a]}$6:${qx[a]}$85"
        x = f"'Quarterly returns'!${qx['World']}$6:${qx['World']}$85"
        rows.append([a, f"=SLOPE({y},{x})", f"=RSQ({y},{x})", f"=COUNT({y})", used[a], lit[a]])
    box(f, top6c, 1, "6c. Betas to world equities (quarterly excess returns since 2006)",
        ["Asset class", "Beta", "R-squared", "Quarters", "Used?", "Literature value"], rows, GREY,
        fmts=[None, "0.00", "0.00", "0", None, "0.00"])

    # ---- 7. Blend ----
    floors = [0.01, 0.01, 0.015, 0.015, 0.03, 0.02, 0.02, 0.03, 0.003, 0.005, 0.003]
    va = {"Australian Equities": (f"G{RE['Australian Equities']}", f"I{RE['Australian Equities']}"),
          "World Equities": (f"G{RE['World Equities']}", f"I{RE['World Equities']}"),
          "Emerging Markets": (f"G{RE['Emerging Markets']}", f"I{RE['Emerging Markets']}"),
          "Listed Property": (f"G{RE['Listed Property']}", f"I{RE['Listed Property']}"),
          "Commodities": (f"F{R['Commodities']}", None),
          "Direct Property": (f"F{R['Direct Property']}", f"G{R['Direct Property']}"),
          "Hedge Funds": (f"F{R['Hedge Funds']}", None),
          "Private Equity": (f"F{R['Private Equity']}", None),
          "Australian Fixed Income": (f"F{R['Australian Fixed Income']}", None),
          "World Fixed Income": (f"F{R['World Fixed Income']}", None),
          "Cash": (f"F{R['Cash']}", f"G{R['Cash']}")}
    rows = []
    for k, a in enumerate(ASSETS):
        rr = RS[a]
        A, B = va[a]
        s1 = 6 + k
        addback = f"=F{s1}" if a in ("Hedge Funds", "Private Equity") else 0
        rows.append([a, f"={A}", f"={B}" if B else None, f"=AVERAGE(B{rr}:C{rr})", addback, f"=H{s1}",
                     f"=D{rr}+E{rr}+F{rr}^2/2", floors[k],
                     f"=IF(COUNT(B{rr}:C{rr})>1,MAX(STDEV(B{rr}:C{rr}),H{rr}),H{rr})",
                     f"=D{s1}", f"=C{s1}", f"=G{s1}", f"=G{s1}-F{rr}^2/2"])
    last7 = box(f, top7, 1, "7. Blend: compound views converted to arithmetic, then combined with the "
                "equilibrium prior (feeds section 1)",
                ["Asset class", "View A (compound)", "View B (compound)", "Composite (compound)",
                 "Fee add-back (views net of fees)", "Volatility", "Composite (gross, arithmetic)",
                 "Minimum view uncertainty", "View dispersion used", "Confidence", "Equilibrium prior",
                 "Net forecast (arithmetic)", "Net forecast (compound)"], rows, PURPLE,
                fmts=[None, PCT2, PCT2, PCT2, PCT2, PCT1, PCT2, PCT2, PCT2, "0.00", PCT2, PCT2, PCT2],
                inputs={(None, 7)}, header_height=54)
    # section 1 now reads the theory-based composite and dispersion
    for k, a in enumerate(ASSETS):
        f[f"B{6 + k}"] = f"=G{RS[a]}"
        f[f"I{6 + k}"] = f"=I{RS[a]}"
    f["B5"] = "Theory-based composite (gross, arithmetic)"
    f["I5"] = "View dispersion"
    f["G18"] = f"=MAX(M{S0}:M{S0 + 10})"
    f["A18"] = "Best single-asset forecast (compound)"
    f["A53"] = "4. Building blocks (v3.4 method, kept for comparison; no longer feeds section 1)"
    comp = {a: f"Forecast!$M${RS[a]}" for a in ASSETS}

    # ---- 8. Uncertainty ----
    top8 = last7 + 3
    rows = []
    for a in ASSETS:
        rr = top8 + 2 + len(rows)
        s = RS[a]
        rows.append([a, f"=M{s}", f"=F{s}", f"=C{rr}/SQRT(5)", f"=I{s}", f"=SQRT(D{rr}^2+E{rr}^2)",
                     f"=B{rr}-1.645*F{rr}", f"=B{rr}+1.645*F{rr}"])
    last8 = box(f, top8, 1, "8. Forecast uncertainty: 90% range for the 5-year annualised (compound) return",
                ["Asset class", "Net forecast (compound)", "Volatility", "Sampling error (vol / sqrt 5)",
                 "View dispersion", "Total standard error", "5th percentile", "95th percentile"],
                rows, ORANGE, fmts=[None, PCT2, PCT1, PCT2, PCT2, PCT2, PCT1, PCT1], header_height=40)

    # ---- 9. Comparison ----
    top9 = last8 + 3
    pub = {"Australian Equities": "DMS (2026): equity premium c.3.5% over bills",
           "World Equities": "J.P. Morgan 2026 LTCMA: global equities 7.0% (USD, 10-15 years)",
           "Emerging Markets": "J.P. Morgan 2026 LTCMA: global equities 7.0% (USD)",
           "Listed Property": "Nareit: REIT dividend yield 3.7% (Aug-2026)",
           "Commodities": "Levine et al. (2018): 3.3% geometric excess, 1877-2015",
           "Direct Property": "Cotality: gross rental yield 3.79% (Aug-2026)",
           "Hedge Funds": "Dichev & Yu (2011): investor returns 3-7% below fund returns",
           "Private Equity": "J.P. Morgan 2026 LTCMA: private equity 10.3% (USD)",
           "Australian Fixed Income": "Vanguard VAF: yield to maturity 5.10%",
           "World Fixed Income": "Vanguard VIF: yield to maturity 5.15% (31-Mar-26)",
           "Cash": "RBA cash rate 4.60% (29-Sep-26)"}
    rows = []
    for k, a in enumerate(ASSETS):
        rr = top9 + 2 + k
        rows.append([a, V34_NET[k], f"=L{RS[a]}", f"=M{RS[a]}", f"=C{rr}-B{rr}", f"=B{25 + k}", pub[a]]
                    + [None] * 2)
    last9 = box(f, top9, 1, "9. Comparison: v3.4 forecasts, 20-year history and published assumptions",
                ["Asset class", "v3.4 net (old method)", "v3.5 net (arithmetic)", "v3.5 net (compound)",
                 "Change in optimiser input", "20-year history (gross, compound)", "Published reference"]
                + [None] * 2, rows, GOLD, fmts=[None, PCT2, PCT2, PCT2, PCT2, PCT2, None],
                header_height=40, notes_col=6)
    for r in range(top9 + 1, last9 + 1):
        merge_text(f, r, 7, 9)

    # ---- 10. Rationale ----
    top10 = last9 + 3
    why = {
        "Australian Equities": ("Grinold-Kroner: forward yield less dilution, plus real earnings growth and inflation, with half the P/E gap to its long-run average closing; averaged with cash plus the DMS premium.",
                                "Valuations now sit above their long-run average, so part of the re-rating behind the 20-year return is assumed to reverse."),
        "World Equities": ("Grinold-Kroner in AUD (relative PPP adds Australian rather than foreign inflation); low dividend yield and a P/E above average; averaged with the build-up.",
                           "History was lifted by a US re-rating and a falling AUD; PPP implies only a small currency drift from here."),
        "Emerging Markets": ("Grinold-Kroner with near-average valuation and higher growth, offset by heavier dilution.",
                             "Little repricing drag because valuations are close to average."),
        "Listed Property": ("REIT dividend yield less issuance plus rent growth; build-up uses the data beta to world equities.",
                            "Yield-based view replaces a history that included the GFC collapse in REITs."),
        "Commodities": ("Collateral (US cash) plus a conservative long-run excess return and the PPP drift.",
                        "Futures returns depend on the shape of the curve (Erb & Harvey 2006), so 20-year history is a poor guide."),
        "Direct Property": ("Cap-rate model: gross rent yield less running costs and depreciation (RBA estimates), plus rent growth; averaged with a beta build-up.",
                            "Net yield is low after costs, so the forecast sits below the appraisal-smoothed history."),
        "Hedge Funds": ("Cash plus a 0.35 equity beta times the expected premium; zero net alpha; fees added back to state it gross.",
                        "Investor-weighted evidence (Dichev & Yu 2011) argues against extrapolating index alpha."),
        "Private Equity": ("Cash plus a 1.2 levered beta times the premium, plus a 1% net premium; fees added back.",
                           "Recent vintages have earned smaller premiums than the 1984-2008 funds in Harris, Jenkinson & Kaplan (2014)."),
        "Australian Fixed Income": ("Starting yield to maturity less expected credit losses; horizon is close to duration, where yield is the best predictor.",
                                    "Yields are far above the 20-year average return, which was depressed by the 2020-22 rate rise."),
        "World Fixed Income": ("Covered interest parity: foreign yields plus the hedge carry from higher AUD cash rates, less a duration and basis adjustment.",
                               "Higher global yields and positive AUD hedge carry lift the forecast above history."),
        "Cash": ("Average of the market-implied path (5-year yield less term premium) and today's cash rate.",
                 "Cash rates are well above the 20-year average after the 2022-26 tightening cycle."),
    }
    rows = [[a, why[a][0]] + [None] * 4 + [why[a][1]] + [None] * 2 for a in ASSETS]
    last10 = box(f, top10, 1, "10. Forecast rationale (for the report)",
                 ["Asset class", "Method and main drivers"] + [None] * 4 + ["Why it differs from history"]
                 + [None] * 2, rows, OLIVE, notes_col=None)
    for r in range(top10 + 1, last10 + 1):
        merge_text(f, r, 2, 6, 40 if r > top10 + 1 else None)
        merge_text(f, r, 7, 9)
        f.cell(r, 1).alignment = Alignment(vertical="center")
    return {"S0": S0, "comp": comp, "RB": RB, "top5": top, "last5": last5, "K": K}


# ---------------------------------------------------------------------------------------------
def downstream(wb, ref):
    """Compare compound (not arithmetic) returns with the target wherever a target test is made."""
    s = wb["Summary"]
    s["A24"] = "Expected return, compound (p.a.)"
    for c, r in zip("BCD", (56, 57, 58)):
        s[f"{c}24"] = f"=Risk!N{r}-Risk!M{r}^2/2"
    s["J6"] = "=Risk!N58-Risk!M58^2/2"
    s["H6"] = "(a) Return: bond yield + 4% (compound)"
    p = wb["Portfolio"]
    for c in "BCDEFG":
        p[f"{c}59"] = f'=IF({c}55-{c}56^2/2>=Inputs!$B$25,"YES","NO")'
    p["A59"] = "Reaches return target? (compound)"
    # every return comparison uses compound returns, like the target test
    comp = lambda r: f"(Risk!$N${r}-Risk!$M${r}^2/2)"
    ck = wb["Checks"]
    ck["A30"] = "Recommended compound return >= current"
    ck["B30"] = f"={comp(58)}-{comp(56)}"
    p["A78"] = "Expected compound return advantage vs current (p.a.)"
    p["D78"] = f"={comp(58)}-{comp(56)}"
    p["C106"] = "Expected return (pre-tax, compound)"
    for i, r in enumerate((108, 109, 110)):
        p[f"C{r}"] = f"={comp(56 + i)}"
    v = wb["Validation"]
    v["A66"] = "Base advantage vs current (compound)"
    v["B66"] = f"={comp(58)}-{comp(56)}"
    rf = wb["Report figures"]
    for k in range(11):
        rf[f"O{308 + k}"] = f"=Forecast!M{ref['S0'] + k}"
    rf["O307"] = "Net forecast (5-year, compound)"


def checks(wb, ref):
    ws = wb["Checks"]
    last = 44
    while ws.cell(last + 1, 1).value not in (None, ""):
        last += 1
    existing = []
    for r in range(7, last + 1):
        existing.append(([ws.cell(r, c).value for c in range(1, 5)],
                         [ws.cell(r, c).number_format for c in range(1, 5)]))
    S0 = ref["S0"]
    rb0 = min(ref["RB"].values())
    new = [
        (["Theory-based forecasts complete (11 asset classes)", f"=COUNT(Forecast!M{S0}:M{S0 + 10})", 11,
          '=IF(B{r}=C{r},"PASS","FAIL")'], "0"),
        (["Compound forecasts below arithmetic for every asset",
          f"=SUMPRODUCT((Forecast!M{S0}:M{S0 + 10}>Forecast!L{S0}:L{S0 + 10})*1)", 0,
          '=IF(B{r}=C{r},"PASS","FAIL")'], "0"),
        (["World bond country weights sum to 100%",
          f"=SUM(Forecast!{ref['K']['w_us'].replace('$', '')}:{ref['K']['w_eu'].replace('$', '')})", 1,
          '=IF(ABS(B{r}-C{r})<0.000001,"PASS","FAIL")'], "0.0%"),
        (["Beta regressions have 30+ quarters", f"=MIN(Forecast!D{rb0}:D{rb0 + 3})", 30,
          '=IF(B{r}>=C{r},"PASS","FAIL")'], "0"),
    ]
    rows, fmts = [], {}
    for i, (vals, nf) in enumerate(existing):
        rows.append(vals)
        fmts[i] = nf
    for vals, nf in new:
        r = 7 + len(rows)
        rows.append([v.replace("{r}", str(r)) if isinstance(v, str) else v for v in vals])
        fmts[len(rows) - 1] = [None, nf, nf, None]
    # clear the old footnotes below the box
    for r in range(last + 1, last + 6):
        ws.cell(r, 1).value = None
    newlast = box(ws, 5, 1, "Tests", ["Test", "Value", "Limit", "Status"], rows, GREEN,
                  fmts=lambda ri, ci: fmts[ri][ci])
    for r in range(7, newlast + 1):
        ws[f"D{r}"].alignment = Alignment(horizontal="center")
    ws["B4"] = (f'=IF(COUNTIF(D7:D{newlast},"FAIL")=0,"ALL CHECKS PASS","FAIL: "&COUNTIF(D7:D{newlast},"FAIL")'
                f'&" check(s)")')
    status_cf(ws, f"D{last + 1}:D{newlast}")
    ws[f"A{newlast + 2}"] = ("REVIEW marks a judgement to disclose in the report, not a model error. "
                             "The master flag counts FAIL only.")
    ws[f"A{newlast + 2}"].font = ws["A7"].font.copy(size=9, italic=True, color="FF595959")
    ws[f"A{newlast + 3}"] = f'=COUNTIF(D7:D{newlast},"REVIEW")&" item(s) flagged REVIEW"'
    ws[f"A{newlast + 3}"].font = ws["A7"].font.copy(size=9, bold=True, color="FF806000")
    # Summary count of REVIEW items follows the new range
    s = wb["Summary"]
    for r in range(70, 90):
        v = s.cell(r, 2).value
        if isinstance(v, str) and "Checks!D7:D" in v and "REVIEW" in v:
            s.cell(r, 2).value = f'=COUNTIF(Checks!D7:D{newlast},"REVIEW")'


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    ref = run(wb)
    downstream(wb, ref)
    checks(wb, ref)
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
