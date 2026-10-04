"""Build FINM3008_EPPIB_Model_v3.4.xlsx from v3.3.

Usage: python tools/build_v34.py <v3.3 input.xlsx> <output.xlsx>

All additions follow the workbook's existing boxed-section layout and palette (see style.py).
Nothing existing is moved: new sections are appended below or beside existing content so that
chart anchors and cross-sheet references stay valid. Charts are restored byte-for-byte from the
source file afterwards by tools/restore_charts.py.
"""
import datetime as dt
import sys

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from style import (BLUE, BROWN, GOLD, GREEN, GREY, NAVY, OLIVE, ORANGE, PURPLE, RED, TEAL,
                   box, font, sheet_title, status_cf)

ASSETS = ["Australian Equities", "World Equities", "Emerging Markets", "Listed Property",
          "Commodities", "Direct Property", "Hedge Funds", "Private Equity",
          "Australian Fixed Income", "World Fixed Income", "Cash"]
PCT1, PCT2 = "0.0%", "0.00%"


def col_range(c0, n):
    return [L(c0 + i) for i in range(n)]


# ---------------------------------------------------------------------------------------------
# 1. Corrections to existing formulas
# ---------------------------------------------------------------------------------------------
def maxgap(prefix, r):
    """Largest absolute weight gap to the recommended row, written cell by cell so it evaluates
    identically in every Excel version (no array evaluation needed)."""
    terms = ",".join(f"ABS({prefix}{c}{r}-{prefix}${c}$58)" for c in "BCDEFGHIJKL")
    return f"=MAX({terms})"


def fix_existing(wb):
    # SUMPRODUCT(MAX(ABS(range-range))) does not array-evaluate MAX in Excel: it returned
    # #VALUE! (Risk P59:P62) or an implicit-intersection value (Validation E61 showed 2.46%,
    # the true largest gap is 13.0%). Replaced with an explicit MAX of the eleven gaps.
    risk = wb["Risk"]
    for r in range(59, 63):
        risk[f"P{r}"] = maxgap("", r)
    wb["Validation"]["E61"] = maxgap("Risk!", 56)

    # Historical portfolio returns were gross of fees although the data notice says the index
    # data are gross and fees must be deducted. Deduct each portfolio's own fee quarterly.
    q = wb["Quarterly returns"]
    rows = {"Z": 56, "AA": 57, "AB": 58, "AC": 63, "AD": 64}
    for r in range(6, 86):
        for col, rr in rows.items():
            q[f"{col}{r}"] = (f"=SUMPRODUCT($N{r}:$X{r},Risk!$B${rr}:$L${rr})"
                              f"-SUMPRODUCT(MMULT(Risk!$B${rr}:$L${rr},Inputs!$D$5:$D$15))/4")
    q["Z3"] = "Constant-weight portfolios (net of fees)"

    wb["Track record"]["A3"] = ("1. Rolling 5-year windows, constant weights (quarterly data, net of "
                                "fees; Recommended is an in-sample back-test)")
    wb["Track record"]["L3"] = "2. By horizon (to March 2026, net of fees)"
    wb["Validation"]["A3"] = "1. Historical backtest (annualised, net of fees)"
    wb["Summary"]["H7"] = "(b) Minimise risk: volatility vs current fund"


# ---------------------------------------------------------------------------------------------
# 2. Performance tab: monthly rebalancing, net of fees (brief: Mar-2021 to Mar-2026 and 20 years)
# ---------------------------------------------------------------------------------------------
P0 = 5            # base row (Mar-2006, wealth = 1)
P1 = 245          # last row (Mar-2026)
R5 = 186          # first month of the 5-year window (Apr-2021)
DATA0 = 257       # Data row for Mar-2006


def build_performance(wb):
    ws = wb.create_sheet("Performance", index=wb.sheetnames.index("Track record") + 1)
    ws.sheet_properties.tabColor = OLIVE[0]
    ws.sheet_view.showGridLines = False
    sheet_title(ws, "Performance")
    widths = {"A": 40, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 12,
              "J": 12, "K": 12, "L": 12, "M": 3}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    # ---- monthly panel (right-hand side, same layout as the Oil panel) ----
    c0 = 14  # column N
    heads = (["Month end"] + ASSETS +
             ["Current gross", "Current net", "Benchmark gross", "Benchmark net", "Active (net)",
              "Current wealth", "Benchmark wealth", "Current drawdown", "Benchmark drawdown",
              "Current drawdown (5y window)", "Benchmark drawdown (5y window)",
              "Current 12m", "Benchmark 12m",
              "Rolling 60m current", "Rolling 60m benchmark", "Rolling 60m active",
              "Rolling 60m TE", "Rolling 60m IR"])
    C = {h: L(c0 + i) for i, h in enumerate(heads)}
    a0, a1 = C[ASSETS[0]], C[ASSETS[-1]]
    rows = []
    for i, r in enumerate(range(P0, P1 + 1)):
        d = DATA0 + i
        row = [f"=Data!A{d}"]
        if r == P0:
            row += [None] * 16 + [1, 1] + [0, 0] + [None] * 9
        else:
            row += [f"=Data!{L(2 + k)}{d}/Data!{L(2 + k)}{d - 1}-1" for k in range(11)]
            row += [f"=SUMPRODUCT({a0}{r}:{a1}{r},Risk!$B$56:$L$56)",
                    f"={C['Current gross']}{r}-$B$8/12",
                    f"=SUMPRODUCT({a0}{r}:{a1}{r},Risk!$B$57:$L$57)",
                    f"={C['Benchmark gross']}{r}-$B$9/12",
                    f"={C['Current net']}{r}-{C['Benchmark net']}{r}",
                    f"={C['Current wealth']}{r - 1}*(1+{C['Current net']}{r})",
                    f"={C['Benchmark wealth']}{r - 1}*(1+{C['Benchmark net']}{r})",
                    f"={C['Current wealth']}{r}/MAX({C['Current wealth']}${P0}:{C['Current wealth']}{r})-1",
                    f"={C['Benchmark wealth']}{r}/MAX({C['Benchmark wealth']}${P0}:{C['Benchmark wealth']}{r})-1"]
            if r >= R5:
                row += [f"={C['Current wealth']}{r}/MAX({C['Current wealth']}${R5 - 1}:{C['Current wealth']}{r})-1",
                        f"={C['Benchmark wealth']}{r}/MAX({C['Benchmark wealth']}${R5 - 1}:{C['Benchmark wealth']}{r})-1"]
            else:
                row += ["", ""]
            if r >= P0 + 12:
                row += [f"={C['Current wealth']}{r}/{C['Current wealth']}{r - 12}-1",
                        f"={C['Benchmark wealth']}{r}/{C['Benchmark wealth']}{r - 12}-1"]
            else:
                row += ["", ""]
            if r >= P0 + 60:
                cw, bw = C["Current wealth"], C["Benchmark wealth"]
                row += [f"=({cw}{r}/{cw}{r - 60})^(12/60)-1",
                        f"=({bw}{r}/{bw}{r - 60})^(12/60)-1",
                        f"={C['Rolling 60m current']}{r}-{C['Rolling 60m benchmark']}{r}",
                        f"=STDEV({C['Active (net)']}{r - 59}:{C['Active (net)']}{r})*SQRT(12)",
                        f"={C['Rolling 60m active']}{r}/{C['Rolling 60m TE']}{r}"]
            else:
                row += ["", "", "", "", ""]
        rows.append(row)

    def pfmt(ri, ci):
        if ci == 0:
            return "mmm-yy"
        h = heads[ci]
        if "wealth" in h:
            return "0.000"
        if h == "Rolling 60m IR":
            return "0.00"
        return PCT2

    box(ws, 3, c0, "Monthly panel, Mar-2006 to Mar-2026 (Data tab index levels; monthly "
        "rebalancing; fees deducted monthly)", heads, rows, GREY, fmts=pfmt, header_height=40)
    for i in range(1, len(heads)):
        ws.column_dimensions[L(c0 + i)].width = 11
    ws.column_dimensions[L(c0)].width = 10
    ws.freeze_panes = "A5"

    # ---- 1. Settings ----
    box(ws, 3, 1, "1. Settings", ["Parameter", "Value", "Source"], [
        ["Panel start (base month)", f"={C['Month end']}{P0}", "Data tab"],
        ["Five-year window starts", f"={C['Month end']}{R5 - 1}", "Brief: Mar-2021"],
        ["Panel end", f"={C['Month end']}{P1}", "Brief: Mar-2026"],
        ["Annual fee: current fund", "=SUMPRODUCT(Portfolio!C6:C16,Inputs!$D$5:$D$15)", "Fee schedule"],
        ["Annual fee: benchmark", "=SUMPRODUCT(Portfolio!D6:D16,Inputs!$D$5:$D$15)", "Fee schedule"],
        ["5-year bond yield at Mar-2021", "=INDEX('RBA 5Y yield'!$B$5:$B$661,MATCH(B6,'RBA 5Y yield'!$A$5:$A$661,1))", "RBA F2"],
        ["5-year bond yield at Mar-2006", "=INDEX('RBA 5Y yield'!$B$5:$B$661,MATCH(B5,'RBA 5Y yield'!$A$5:$A$661,1))", "RBA F2"],
        ["Rebalancing", "Monthly", "Brief"],
    ], GOLD, fmts=lambda ri, ci: ("mmm-yy" if ri < 3 else PCT2) if ci == 1 else None)
    return ws, C


def finish_performance(ws, C):
    """Write sections 2 to 4 (left-hand side) once panel columns are known."""
    cg, cn, bg, bn, act = (C["Current gross"], C["Current net"], C["Benchmark gross"],
                           C["Benchmark net"], C["Active (net)"])
    cw, bw = C["Current wealth"], C["Benchmark wealth"]
    cash = C["Cash"]

    def rng(col, a, b):
        return f"{col}{a}:{col}{b}"

    def ann(col, a, b):
        n = b - a + 1
        return f"EXP(SUMPRODUCT(LN(1+{rng(col, a, b)})))^(12/{n})-1"

    def metrics(a, b, dd_c, dd_b, worst_from, yld):
        n = b - a + 1
        rc = f"({cw}{b}/{cw}{a - 1})^(12/{n})-1"
        rb = f"({bw}{b}/{bw}{a - 1})^(12/{n})-1"
        cashr = ann(cash, a, b)
        te = f"STDEV({rng(act, a, b)})*SQRT(12)"
        out = [
            [f"={rc}", f"={rb}"],
            [f"={ann(cg, a, b)}", f"={ann(bg, a, b)}"],
            [f"=STDEV({rng(cn, a, b)})*SQRT(12)", f"=STDEV({rng(bn, a, b)})*SQRT(12)"],
            [f"=(({rc})-({cashr}))/(STDEV({rng(cn, a, b)})*SQRT(12))",
             f"=(({rb})-({cashr}))/(STDEV({rng(bn, a, b)})*SQRT(12))"],
            [f"={te}", None],
            [f"=(({rc})-({rb}))/({te})", None],
            [f"=SLOPE({rng(cn, a, b)},{rng(bn, a, b)})", 1],
            [f"=MIN({rng(dd_c, a, b)})", f"=MIN({rng(dd_b, a, b)})"],
            [f"=MIN({rng(C['Current 12m'], worst_from, b)})", f"=MIN({rng(C['Benchmark 12m'], worst_from, b)})"],
            [f"=AVERAGEIFS({rng(cn, a, b)},{rng(bn, a, b)},\">0\")/AVERAGEIFS({rng(bn, a, b)},{rng(bn, a, b)},\">0\")", 1],
            [f"=AVERAGEIFS({rng(cn, a, b)},{rng(bn, a, b)},\"<0\")/AVERAGEIFS({rng(bn, a, b)},{rng(bn, a, b)},\"<0\")", 1],
            [f"=COUNTIF({rng(act, a, b)},\">0\")/COUNT({rng(act, a, b)})", None],
            [f"={yld}+Inputs!$B$24", f"={yld}+Inputs!$B$24"],
        ]
        return out

    m5 = metrics(R5, P1, C["Current drawdown (5y window)"], C["Benchmark drawdown (5y window)"], R5 + 11, "$B$10")
    m20 = metrics(P0 + 1, P1, C["Current drawdown"], C["Benchmark drawdown"], P0 + 12, "$B$11")
    labels = ["Annualised return, net of fees", "Annualised return, gross of fees",
              "Volatility (annualised)", "Sharpe ratio (excess over bank bills)", "Tracking error",
              "Information ratio", "Beta to benchmark", "Maximum drawdown", "Worst 12-month return",
              "Up-market capture", "Down-market capture", "Months outperforming benchmark",
              "Return target (5-year yield + 4%)"]
    top = 15
    rows = []
    for i, lab in enumerate(labels):
        r = top + 2 + i
        a, b = m5[i]
        c, d = m20[i]
        diff5 = f"=B{r}-C{r}" if b is not None and i not in (6, 9, 10, 12) else None
        diff20 = f"=E{r}-F{r}" if d is not None and i not in (6, 9, 10, 12) else None
        rows.append([lab, a, b, diff5, c, d, diff20])
    rt = top + 2 + len(labels)
    rows.append(["Return target met?", f'=IF(B{top + 2}>=B{rt - 1},"MET","NOT MET")',
                 f'=IF(C{top + 2}>=C{rt - 1},"MET","NOT MET")', None,
                 f'=IF(E{top + 2}>=E{rt - 1},"MET","NOT MET")',
                 f'=IF(F{top + 2}>=F{rt - 1},"MET","NOT MET")', None])

    def mfmt(ri, ci):
        if ci == 0:
            return None
        if ri in (3, 5, 6, 9, 10):
            return "0.00"
        return PCT2
    last = box(ws, top, 1, "2. Performance summary (monthly rebalancing, net of fees)",
               ["Measure", "Current: 5 years", "Benchmark: 5 years", "Difference",
                "Current: 20 years", "Benchmark: 20 years", "Difference"], rows, OLIVE,
               fmts=mfmt, header_height=28)
    status_cf(ws, f"B{last}:F{last}")
    perf_rows = {lab: top + 2 + i for i, lab in enumerate(labels)}
    perf_rows["met"] = last

    # ---- 3. Non-overlapping five-year blocks ----
    top3 = last + 3
    blocks = [(P0 + 1, P0 + 60), (P0 + 61, P0 + 120), (P0 + 121, P0 + 180), (P0 + 181, P1)]
    rows = []
    for a, b in blocks:
        rc = f"({cw}{b}/{cw}{a - 1})^(12/60)-1"
        rb = f"({bw}{b}/{bw}{a - 1})^(12/60)-1"
        te = f"STDEV({rng(act, a, b)})*SQRT(12)"
        cashr = ann(cash, a, b)
        yld = f"INDEX('RBA 5Y yield'!$B$5:$B$661,MATCH({C['Month end']}{a - 1},'RBA 5Y yield'!$A$5:$A$661,1))"
        r = top3 + 2 + len(rows)
        rows.append([f"=TEXT({C['Month end']}{a - 1},\"mmm-yy\")&\" to \"&TEXT({C['Month end']}{b},\"mmm-yy\")",
                     f"={rc}", f"={rb}", f"=B{r}-C{r}", f"={te}", f"=D{r}/E{r}",
                     f"=(B{r}-({cashr}))/(STDEV({rng(cn, a, b)})*SQRT(12))",
                     f"=(C{r}-({cashr}))/(STDEV({rng(bn, a, b)})*SQRT(12))",
                     f"={yld}+Inputs!$B$24", f'=IF(B{r}>=I{r},"MET","NOT MET")'])
    f3, l3 = top3 + 2, top3 + 5
    rows.append(["Change, first block to last", None, None, f"=D{l3}-D{f3}", f"=E{l3}-E{f3}",
                 f"=F{l3}-F{f3}", f"=G{l3}-G{f3}", f"=H{l3}-H{f3}", None, None])
    roll_a, roll_ir = C["Rolling 60m active"], C["Rolling 60m IR"]
    rows.append(["Rolling 60m windows with positive active return",
                 f"=COUNTIF({rng(roll_a, P0 + 60, P1)},\">0\")/COUNT({rng(roll_a, P0 + 60, P1)})"] + [None] * 8)
    half = (P0 + 60 + P1) // 2
    rows.append(["Average rolling 60m active: first half of windows",
                 f"=AVERAGE({rng(roll_a, P0 + 60, half)})"] + [None] * 8)
    rows.append(["Average rolling 60m active: second half of windows",
                 f"=AVERAGE({rng(roll_a, half + 1, P1)})"] + [None] * 8)
    rows.append(["Slope of rolling 60m active return (pp per year)",
                 f"=SLOPE({rng(roll_a, P0 + 60, P1)},{rng(C['Month end'], P0 + 60, P1)})*365.25"] + [None] * 8)

    def f3fmt(ri, ci):
        if ci in (5, 6, 7):
            return "0.00"
        if ci in (1, 2, 3, 4, 8):
            return PCT2
        return None
    last3 = box(ws, top3, 1, "3. Is performance improving? Non-overlapping 5-year blocks (net of fees)",
                ["Block", "Current", "Benchmark", "Active return", "Tracking error",
                 "Information ratio", "Sharpe: current", "Sharpe: benchmark", "Target (yield + 4%)",
                 "Target met?"], rows, PURPLE, fmts=f3fmt, header_height=28)
    status_cf(ws, f"J{f3}:J{l3}")
    ws[f"B{l3 + 2}"].number_format = PCT1

    # ---- 4. Attribution of active return ----
    top4 = last3 + 3
    rows = []
    a5, a20 = (R5, P1), (P0 + 1, P1)
    for k, asset in enumerate(ASSETS):
        col = C[asset]
        r = top4 + 2 + k
        rows.append([asset, f"=Portfolio!C{6 + k}", f"=Portfolio!D{6 + k}", f"=B{r}-C{r}",
                     f"=12*(AVERAGE({rng(col, *a5)})-AVERAGE({rng(bg, *a5)}))", f"=D{r}*E{r}",
                     f"=12*(AVERAGE({rng(col, *a20)})-AVERAGE({rng(bg, *a20)}))", f"=D{r}*G{r}"])
    fa, fl = top4 + 2, top4 + 12
    rf = fl + 1
    rows.append(["Fee effect", None, None, None, None, "=-($B$8-$B$9)", None, "=-($B$8-$B$9)"])
    rows.append(["Total explained (arithmetic)", f"=SUM(B{fa}:B{fl})", f"=SUM(C{fa}:C{fl})",
                 f"=SUM(D{fa}:D{fl})", None, f"=SUM(F{fa}:F{rf})", None, f"=SUM(H{fa}:H{rf})"])
    rows.append(["Active return, arithmetic (12 x mean monthly active)", None, None, None, None,
                 f"=12*AVERAGE({rng(act, *a5)})", None, f"=12*AVERAGE({rng(act, *a20)})"])
    rows.append(["Compounding difference (geometric less arithmetic)", None, None, None, None,
                 f"=D{perf_rows['Annualised return, net of fees']}-F{rf + 2}", None,
                 f"=G{perf_rows['Annualised return, net of fees']}-H{rf + 2}"])

    def f4fmt(ri, ci):
        if ci in (1, 2, 3):
            return PCT1
        if ci > 0:
            return PCT2
        return None
    last4 = box(ws, top4, 1, "4. Attribution of active return (allocation effect; asset classes held "
                "passively)", ["Asset class", "Current weight", "Benchmark weight", "Active weight",
                               "5y: asset less benchmark (p.a.)", "5y: contribution",
                               "20y: asset less benchmark (p.a.)", "20y: contribution"],
                rows, TEAL, fmts=f4fmt, header_height=40, bold_rows=(12,))
    attrib = {"first": fa, "last": fl, "fee": rf, "total": rf + 1, "arith": rf + 2}

    # ---- 5. Notes ----
    top5 = last4 + 3
    notes = [
        ["Rebalancing", "Monthly to the stated weights, as the brief specifies; no transaction costs."],
        ["Fees", "Deducted monthly at one twelfth of each portfolio's annual fee (Fee schedule)."],
        ["Illiquid series", "Direct property and private equity indices update quarterly, so monthly "
                            "benchmark volatility is understated; reported (not unsmoothed) returns are "
                            "used because this tab measures realised performance."],
        ["Attribution", "The fund and benchmark hold each asset class passively, so all active return "
                        "is allocation effect: (w fund - w benchmark) x (r asset - r benchmark), "
                        "plus the fee difference."],
        ["Target", "5-year Australian government bond yield at the start of each window plus 4%."],
    ]
    notes = [n + [None] * 8 for n in notes]
    box(ws, top5, 1, "5. Method notes", ["Item", "Note"] + [None] * 8, notes, GREY)
    for r in range(top5 + 1, top5 + 2 + len(notes) - 1 + 1):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        ws[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if r > top5 + 1:
            ws.row_dimensions[r].height = 27
    return perf_rows, attrib, {"blocks_first": f3, "blocks_last": l3, "trend": l3 + 1,
                               "roll_pos": l3 + 2, "roll_slope": l3 + 5}


def main(src, dst):
    wb = openpyxl.load_workbook(src)
    fix_existing(wb)
    ws, C = build_performance(wb)
    perf = finish_performance(ws, C)
    import extend
    figs = extend.run(wb, C, perf)
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)
    return figs


if __name__ == "__main__":
    import json
    figs = main(sys.argv[1], sys.argv[2])
    json.dump({"fig14": figs[0], "fig15": figs[1]}, open(sys.argv[2] + ".figs.json", "w"))
