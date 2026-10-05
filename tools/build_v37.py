"""v3.7 changes applied on top of the v3.6 build (before re-optimisation).

* Oil floor in the optimiser: the US$150 central-case loss of every forecast-dependent optimised
  set must be no worse than the current fund's (Inputs B50). Constraint cost shown as "No oil
  floor"; a stricter floor at the benchmark's loss is priced alongside (it costs more and pushes
  the robust vs resampled gap above the 3-point stability limit, so it is not the base case).
* Stage 2 overweight cap (Inputs B47) now caps the alternatives' upper bounds in the optimiser.
* Stage 2 scorecard: return scores are formula-driven from the net forecasts.
* Tracking-error setting relabelled as a deliberate buffer below the brief limit.
* The fixed-alternatives comparison case is named for what it is (fixed weights), since the
  scorecard is now only a cross-check.
* Integrity table re-tests the oil floor; Checks and Summary report it.
"""
import sys
from copy import copy

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Alignment, Border, Font, PatternFill

from impl import EQ, te_col
from style import GRID, PURPLE, _side, box, status_cf

OIL = "Oil!$F$28:$F$38"
FLOOR = "Inputs!$B$50"
THIN = _side("thin", GRID)
NOTE = Font(name="Arial", size=9, italic=True, color="FF595959")
RENAME = [
    ("Alternatives fixed at the Stage 2 scorecard weights (Commodities 3%, Direct Property 5%, Private "
     "Equity 1%)", "Alternatives fixed at Commodities 3%, Direct Property 5%, Private Equity 1%; other "
     "weights optimised"),
    ("Alternatives fixed by scorecard", "Alternatives at fixed weights"),
    ("alternatives fixed by scorecard", "alternatives at fixed weights"),
    ("risk parity holds the scorecard alternatives", "risk parity holds the fixed alternatives weights"),
]


def blank(c):
    c.value = None
    c.border = Border()
    c.fill = PatternFill(fill_type=None)
    c.font = Font(name="Arial", size=10)
    c.number_format = "General"
    c.alignment = Alignment()


def add_row(ws, last, c0, c1):
    """Extend a box by one body row below `last` (styles copied, bottom edge moved down)."""
    for c in range(c0, c1 + 1):
        src, dst = ws.cell(last, c), ws.cell(last + 1, c)
        dst._style = copy(src._style)
        b = src.border
        src.border = Border(left=b.left, right=b.right, top=b.top, bottom=THIN)
    return last + 1


def add_col(ws, r0, r1, src_col, new_col):
    """Extend a box by one column to the right (styles copied, right edge moved across)."""
    for r in range(r0, r1 + 1):
        src, dst = ws.cell(r, src_col), ws.cell(r, new_col)
        dst._style = copy(src._style)
        b = src.border
        src.border = Border(left=b.left, right=THIN, top=b.top, bottom=b.bottom)


def rename_all(wb):
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    v = c.value
                    for a, b in RENAME:
                        v = v.replace(a, b)
                    if v != c.value:
                        c.value = v


def inputs(wb):
    ws = wb["Inputs"]
    ws["C45"] = "Judgement: buffer below the 1.5% brief limit"
    ws["C47"] = "Judgement; caps Stage 2 upper bounds"
    add_row(ws, 49, 1, 3)
    ws["A50"] = "Oil floor in optimiser: US$150 central-case loss no worse than the current fund"
    ws["B50"] = "=Oil!$F$43"
    ws["C50"] = "Policy (Oil section 4)"
    ws["B50"].number_format = "0.00%"
    ws["B50"].font = Font(name="Arial", size=10, color="FF000000")


def portfolio(wb):
    ws = wb["Portfolio"]
    for k, r in enumerate(range(10, 14)):        # Stage 2 upper bound: benchmark + overweight cap
        ws[f"G{r}"] = f"=MIN(Inputs!F{9 + k},Inputs!C{9 + k}+Inputs!$B$47)"
    # scorecard: return score from the net (arithmetic) forecasts, scaled across the 11 classes
    for k, r in enumerate(range(23, 27)):
        ws[f"B{r}"] = (f"=ROUND(1+4*(K{10 + k}-MIN($K$6:$K$16))/(MAX($K$6:$K$16)-MIN($K$6:$K$16)),0)")
        ws[f"B{r}"].font = copy(ws[f"G{r}"].font)
        ws[f"B{r}"].fill = copy(ws[f"G{r}"].fill)
    ws["A28"] = ("Return score is formula-driven: net forecast (arithmetic) scaled from 1 (lowest of the 11 "
                 "classes) to 5 (highest). The other four scores are judgement.")
    ws["A28"].font = NOTE
    # Solver constraints: oil floor row
    add_row(ws, 38, 1, 5)
    ws["A39"] = "US$150 oil central case (policy floor)"
    ws["B39"] = f"=SUMPRODUCT(E6:E16,{OIL})"
    ws["C39"] = f"={FLOOR}"
    ws["D39"] = None
    ws["E39"] = '=IF(B39>=C39-0.00001,"MET","NOT MET")'
    for c in "BC":
        ws[f"{c}39"].number_format = "0.00%"
    status_cf(ws, "E39")
    # constraint cost: add "No oil floor"
    add_col(ws, 41, 59, 7, 8)
    add_col(ws, 41, 59, 8, 9)
    ws["H42"] = "No oil floor"
    ws["I42"] = "Oil floor at benchmark loss"
    for c in "HI":
        for r in range(43, 54):
            ws[f"{c}{r}"] = 0
        for r in range(55, 60):
            ws[f"{c}{r}"] = ws[f"G{r}"].value.replace("G43:G53", f"{c}43:{c}53").replace("G55", f"{c}55") \
                .replace("G56", f"{c}56")
    # scope descriptions: the oil floor applies to every scope portfolio unless reopt says otherwise
    ws["A125"] = ("9. Asset-class scope: portfolios with and without some asset classes (recommendation "
                  "unchanged; oil floor applied to all)")


def report_figures(wb):
    ws = wb["Report figures"]
    add_row(ws, 183, 14, 16)
    add_row(ws, 184, 14, 16)
    for r, c in ((184, "H"), (185, "I")):
        ws[f"N{r}"] = f"=Portfolio!{c}42"
        ws[f"O{r}"] = f"=Portfolio!{c}55-Portfolio!{c}56^2/2"
        ws[f"P{r}"] = "=Inputs!$B$25"


def validation(wb):
    """Rebuild the integrity table (section 7) with the oil floor re-tested."""
    ws = wb["Validation"]
    top = 85
    for r in range(top, 108):
        for c in range(1, 10):
            blank(ws.cell(r, c))
        ws.row_dimensions[r].height = None
    keep = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        if str(cf.sqref) != "H87:H104":
            for rule in cf.rules:
                keep.add(str(cf.sqref), rule)
    ws.conditional_formatting = keep

    m, p = wb["Methods"], wb["Portfolio"]
    sets = []   # (name, range, equity band, TE limit, oil floor applies)
    for c in "DEFGH":
        sets.append((f"Methods: {m[f'{c}5'].value}", f"Methods!{c}6:{c}16", True, "Inputs!$B$45", c != "H"))
    for c, te in zip("BCDE", [None, "Inputs!$B$45", 0.0075, "Inputs!$B$27"]):
        sets.append((f"Sensitivity: {ws[f'{c}29'].value}", f"Validation!{c}30:{c}40", True, te, c != "B"))
    for c in "BC":
        sets.append((f"Out-of-sample: {ws[f'{c}16'].value}", f"Validation!{c}17:{c}27", True, None, False))
    for c, te in zip("CDEFGHI", ["Inputs!$B$27", None, "Inputs!$B$45", "Inputs!$B$45", None, "Inputs!$B$45",
                                 "Inputs!$B$45"]):
        sets.append((f"Constraint cost: {p[f'{c}42'].value}", f"Portfolio!{c}43:{c}53", c not in "EG", te,
                     c not in "GH"))
    rows = []
    for name, rng, eqband, te, oil in sets:
        r = top + 2 + len(rows)
        b = rng.split("!")[1].split(":")[1]
        rows.append([name, f"=SUM({rng})", f"=MIN({rng})", f"={rng.split('!')[0]}!{b}",
                     f"=SUMPRODUCT({rng},{EQ})", f"={te_col(rng)}",
                     (f"={te}" if isinstance(te, str) else te) if te is not None else "n/a",
                     f"=SUMPRODUCT({rng},{OIL})",
                     f'=IF(AND(ABS(B{r}-1)<0.0005,C{r}>=-0.000001,D{r}>=Inputs!$B$34-0.00001,'
                     + (f'E{r}>=Inputs!$B$31-0.00001,E{r}<=Inputs!$B$32+0.0005,' if eqband else '')
                     + (f'H{r}>={FLOOR}-0.00002,' if oil else '')
                     + (f'F{r}<=G{r}+0.0002' if te is not None else 'TRUE') + '),"PASS","REVIEW")'])
    for title, r0, te in (("Frontier: no TE limit (10 points)", 113, None),
                          ("Frontier: with TE limit (10 points)", 131, "Inputs!$B$45")):
        cols = [chr(ord("B") + i) for i in range(10)]
        r = top + 2 + len(rows)
        sums = ",".join(f"ABS(SUM({c}{r0}:{c}{r0 + 10})-1)" for c in cols)
        cash = ",".join(f"{c}{r0 + 10}" for c in cols)
        eqs = ",".join(f"SUMPRODUCT({c}{r0}:{c}{r0 + 10},{EQ})" for c in cols)
        tes = ",".join(te_col(f"{c}{r0}:{c}{r0 + 10}") for c in cols)
        oils = ",".join(f"SUMPRODUCT({c}{r0}:{c}{r0 + 10},{OIL})" for c in cols)
        rows.append([title, f"=1+MAX({sums})", f"=MIN(B{r0}:K{r0 + 10})", f"=MIN({cash})",
                     f"=MAX({eqs})", f"=MAX({tes})", f"={te}" if te else "n/a", f"=MIN({oils})",
                     f'=IF(AND(ABS(B{r}-1)<0.0005,C{r}>=-0.000001,D{r}>=Inputs!$B$34-0.00001,'
                     f'E{r}<=Inputs!$B$32+0.0005,H{r}>={FLOOR}-0.00002,'
                     + (f'F{r}<=G{r}+0.0002' if te else 'TRUE') + '),"PASS","REVIEW")'])

    def fv(ri, ci):
        return {1: "0.0000", 2: "0.00%", 3: "0.0%", 4: "0.0%", 5: "0.00%", 6: "0.00%", 7: "0.00%"}.get(ci)
    last = box(ws, top, 1, "7. Integrity of pasted Solver outputs (re-tested against current inputs)",
               ["Weight set", "Sum of weights", "Smallest weight", "Cash", "Equity (worst point)",
                "Tracking error (worst point)", "TE limit", "US$150 oil (worst point)", "Status"], rows,
               PURPLE, fmts=fv, header_height=40)
    status_cf(ws, f"I{top + 2}:I{last}")
    ws[f"A{last + 1}"] = ("Pasted Solver outputs: re-run Solver for any REVIEW row after an input change. "
                          "Oil floor tested where it applies.")
    ws[f"A{last + 1}"].font = NOTE
    assert last + 1 < 111, last
    ws["A111"] = "8. Efficient frontier: brief constraints and oil floor, no TE limit"
    ws["A129"] = "9. Efficient frontier: brief constraints and oil floor, with TE limit"
    return f"Validation!$I${top + 2}:$I${last}"


def checks(wb, integ):
    ws = wb["Checks"]
    ws["B28"] = "=SUM(Portfolio!B43:I53)"
    ws["C28"] = "=8"
    ws["B39"] = f'=COUNTIF({integ},"REVIEW")'
    add_row(ws, 49, 1, 4)
    ws["A50"] = "Recommended US$150 oil loss no worse than the floor"
    ws["B50"] = "=Portfolio!B39"
    ws["C50"] = f"={FLOOR}"
    ws["D50"] = '=IF(B50>=C50-0.00001,"PASS","FAIL")'
    for c in "BC":
        ws[f"{c}50"].number_format = "0.00%"
    status_cf(ws, "D50")
    ws["B4"] = '=IF(COUNTIF(D7:D50,"FAIL")=0,"ALL CHECKS PASS","FAIL: "&COUNTIF(D7:D50,"FAIL")&" check(s)")'
    ws["A52"] = None
    ws["A53"] = '=COUNTIF(D7:D50,"REVIEW")&" item(s) flagged REVIEW"'
    ws["A53"].font = copy(ws["A51"].font)
    ws["A52"] = ws["A51"].value
    ws["A52"].font = copy(ws["A51"].font)
    ws["A51"] = None


def summary(wb):
    ws = wb["Summary"]
    add_row(ws, 16, 8, 11)
    ws["H17"] = "(f) Policy: US$150 oil loss no worse than current fund"
    ws["I17"] = f"={FLOOR}"
    ws["J17"] = "=Oil!$F$45"
    ws["K17"] = '=IF(J17>=I17-0.00001,"MET","NOT MET")'
    status_cf(ws, "K17")
    ws["H76"] = "Oil floor (current fund): return cost, compound"
    ws["K76"] = "=(Portfolio!$H$55-Portfolio!$H$56^2/2)-(Portfolio!$B$55-Portfolio!$B$56^2/2)"
    ws["H77"] = "Stricter floor (benchmark): extra return cost"
    ws["K77"] = "=(Portfolio!$B$55-Portfolio!$B$56^2/2)-(Portfolio!$I$55-Portfolio!$I$56^2/2)"
    add_row(ws, 77, 8, 11)
    ws["H78"] = "Stricter floor (benchmark): oil loss"
    ws["K78"] = "=SUMPRODUCT(Portfolio!$I$43:$I$53,Oil!$F$28:$F$38)"
    for c in ("K76", "K77", "K78"):
        ws[c].number_format = "0.00%"


def misc(wb):
    wb["Oil"]["A150"] = "Recommended (oil floor built in)"
    wb["Methods"]["A3"] = ("1. Method weights (alternatives optimised within the Stage 2 cap; Australian fixed "
                           "income floor 12%; oil floor; risk parity holds the fixed alternatives weights)")


def trim(wb, notes):
    """Move the discussion-only text out of the model (it belongs in the report)."""
    ws = wb["Oil"]
    top = 178
    lines = ["# Oil: other mitigation tools (removed from the model; rewrite in your own words for the report)",
             "", "| Tool | How it helps | Cost or drawback |", "|---|---|---|"]
    for r in range(top + 2, top + 8):
        lines.append(f"| {ws[f'A{r}'].value} | {ws[f'B{r}'].value} | {ws[f'F{r}'].value} |")
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= top:
            ws.unmerge_cells(str(rng))
    for r in range(top, top + 9):
        for c in range(1, 10):
            blank(ws.cell(r, c))
        ws.row_dimensions[r].height = None
    with open(notes, "w") as fh:
        fh.write("\n".join(lines) + "\n")
    p = wb["Performance"]
    p["B69"] = "Indices update quarterly: reported returns used, so monthly volatility is understated."
    p["B70"] = "Passive holdings: all active return is allocation, (w fund - w benchmark) x (r asset - r benchmark), plus fees."
    wb["Portfolio"]["A28"] = ("Return score: net arithmetic forecast scaled from 1 (lowest of the 11 classes) to 5 "
                              "(highest). Other scores are judgement.")


def main(src, dst, notes):
    wb = openpyxl.load_workbook(src)
    rename_all(wb)
    inputs(wb)
    portfolio(wb)
    report_figures(wb)
    integ = validation(wb)
    checks(wb, integ)
    summary(wb)
    misc(wb)
    trim(wb, notes)
    wb.calculation.fullCalcOnLoad = True
    wb.save(dst)


if __name__ == "__main__":
    main(*sys.argv[1:4])
