"""Asset-class scope: alternative portfolios with and without some asset classes (Portfolio section 9).

Every portfolio is optimised under the same forecasts, covariance and brief constraints as the
recommendation (weights pasted by reopt.py) and scored on the same measures, including the
bootstrap. The recommendation itself is unchanged.
"""
from copy import copy

from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from style import GREEN, NAVY, TEAL, box, status_cf

ASSETS = ["Australian Equities", "World Equities", "Emerging Markets", "Listed Property",
          "Commodities", "Direct Property", "Hedge Funds", "Private Equity",
          "Australian Fixed Income", "World Fixed Income", "Cash"]
PCT1, PCT2 = "0.0%", "0.00%"
NAMES = ["Recommended", "Alternatives optimised", "No alternatives", "Alternatives at benchmark weights",
         "Core five only"]
DESC = ["Stage 2 alternatives fixed by judgement (Commodities 3%, Direct Property 5%, Private Equity 1%)",
        "Commodities, Direct Property, Hedge Funds and Private Equity chosen by the optimiser (0 to upper bound)",
        "Commodities, Direct Property, Hedge Funds and Private Equity excluded",
        "Alternatives held at benchmark weights (Commodities 2%, Direct Property 5%, Hedge Funds 3%, Private Equity 3%)",
        "Australian and world equities, both bond classes and cash only; cannot reach 1.1% TE, so optimised at the 1.5% brief limit"]
TOP = 125                       # Portfolio row of the weights box
W0, W1 = TOP + 2, TOP + 12      # weight rows (11 assets)
BOOT_ANN = ["AC", "AD", "AE", "AF"]     # Bootstrap annual returns for portfolios 2-5
BOOT_CUM = ["L", "M", "N", "O"]         # Bootstrap 5-year path returns
BOOT_REL = ["P", "Q", "R", "S"]         # Bootstrap log relative to benchmark


def _extend_box(ws, rows, last_old, new_cols):
    """Extend an existing box to the right: interior style for the old edge and new columns,
    edge style for the final new column."""
    for r in rows:
        interior = copy(ws[f"{L(last_old - 1)}{r}"]._style)
        edge = copy(ws[f"{L(last_old)}{r}"]._style)
        ws[f"{L(last_old)}{r}"]._style = copy(interior)
        for i, c in enumerate(new_cols):
            ws[f"{c}{r}"]._style = copy(edge if i == len(new_cols) - 1 else interior)


def bootstrap(wb):
    b = wb["Bootstrap"]
    # annual portfolio returns (section 2), columns AC:AF after AB (Recommended)
    _extend_box(b, range(18, 98), 28, BOOT_ANN)
    for k, c in enumerate(BOOT_ANN):
        pc = "CDEF"[k]
        b[f"{c}20"] = NAMES[k + 1]
        for r in range(21, 98):
            b[f"{c}{r}"] = f"=SUMPRODUCT(MMULT($O{r}:$Y{r},Portfolio!${pc}${W0}:${pc}${W1}))"
        b.column_dimensions[c].width = 12
    # simulated paths (section 3), columns L:S after K
    _extend_box(b, range(99, 1102), 11, BOOT_CUM + BOOT_REL)
    for k in range(4):
        cum, rel, ann = BOOT_CUM[k], BOOT_REL[k], BOOT_ANN[k]
        b[f"{cum}101"] = NAMES[k + 1]
        b[f"{rel}101"] = f"Relative: {NAMES[k + 1].lower()} (log)"
        for r in range(102, 1102):
            terms = "*".join(f"(1+INDEX(${ann}$21:${ann}$97,{d}{r}))" for d in "BCDEF")
            b[f"{cum}{r}"] = f"={terms}-1"
            b[f"{rel}{r}"] = f"=LN((1+{cum}{r})/(1+H{r}))"
        b.column_dimensions[cum].width = 12
        b.column_dimensions[rel].width = 12


def portfolio(wb, oil_sev_row):
    ws = wb["Portfolio"]
    rows = []
    for k, a in enumerate(ASSETS):
        rows.append([a, f"=E{6 + k}", None, None, None, None])
    rows.append(["Total"] + [f"=SUM({c}{W0}:{c}{W1})" for c in "BCDEF"])
    last = box(ws, TOP, 1, "9. Asset-class scope: portfolios with and without some asset classes "
               "(recommendation unchanged)",
               ["Asset class"] + NAMES, rows, NAVY, fmts=[None] + [PCT1] * 5, header_height=40,
               inputs={(ri, ci) for ri in range(11) for ci in range(2, 6)}, bold_rows=(11,))
    # descriptions
    top_d = last + 2
    drows = [[NAMES[k], DESC[k]] + [None] * 4 for k in range(5)]
    lastd = box(ws, top_d, 1, "9b. How each portfolio is built (same forecasts, risk model and constraints)", ["Portfolio", "Construction"] + [None] * 4,
                drows, TEAL)
    for r in range(top_d + 1, lastd + 1):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(r, 2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(r, 1).alignment = Alignment(vertical="center")
        if r > top_d + 1:
            ws.row_dimensions[r].height = 27

    # metrics
    top_m = lastd + 3
    m0 = top_m + 2
    COV = "Risk!$B$41:$L$51"
    BEN = "$D$6:$D$16"
    CUR = "$C$6:$C$16"
    metrics = []

    def W(c):
        return f"{c}${W0}:{c}${W1}"

    labels = ["Expected return (arithmetic, optimiser input)", "Expected return (compound)", "Volatility",
              "Tracking error (parametric)", "Tracking error (bootstrap)",
              "5-year loss probability (bootstrap)", "5-year loss probability (normal)",
              "Sharpe ratio (excess over expected cash)", "Management fee (p.a.)", "Equity weight",
              "Illiquid weight", "Global financial crisis", "COVID shock", "Inflation sell-off (2022)",
              "US$150 oil (central case)", "US$150 oil (severe case, 2008 scaled)", "One-way turnover",
              "Transition cost (A$m)", "Compound advantage vs current, after amortised costs",
              "Brief constraints (TE within 1.5% parametric and bootstrap, equity band, illiquid, cash)"]
    rr = {lab: m0 + i for i, lab in enumerate(labels)}
    for i, lab in enumerate(labels):
        row = [lab]
        for k, c in enumerate("BCDEF"):
            w = W(c)
            if lab.startswith("Expected return (arith"):
                v = f"=SUMPRODUCT({w},Forecast!$G$6:$G$16)"
            elif lab == "Expected return (compound)":
                v = f"={c}{rr[labels[0]]}-{c}{rr['Volatility']}^2/2"
            elif lab == "Volatility":
                v = f"=SQRT(SUMPRODUCT(MMULT({COV},{w}),{w}))"
            elif lab == "Tracking error (parametric)":
                v = f"=SQRT(SUMPRODUCT(MMULT({COV},{w}-{BEN}),{w}-{BEN}))"
            elif lab == "Tracking error (bootstrap)":
                v = "=Bootstrap!$D$12" if k == 0 else f"=STDEV(Bootstrap!{BOOT_REL[k - 1]}102:{BOOT_REL[k - 1]}1101)/SQRT(5)"
            elif lab.startswith("5-year loss probability (boot"):
                v = ("=Bootstrap!$D$6" if k == 0 else
                     f'=COUNTIF(Bootstrap!{BOOT_CUM[k - 1]}102:{BOOT_CUM[k - 1]}1101,"<0")/COUNT(Bootstrap!{BOOT_CUM[k - 1]}102:{BOOT_CUM[k - 1]}1101)')
            elif lab.startswith("5-year loss probability (normal"):
                v = f"=NORMSDIST(-{c}{rr['Expected return (compound)']}*SQRT(5)/{c}{rr['Volatility']})"
            elif lab.startswith("Sharpe"):
                v = f"=({c}{rr[labels[0]]}-Forecast!$G$16)/{c}{rr['Volatility']}"
            elif lab.startswith("Management fee"):
                v = f"=SUMPRODUCT({w},Inputs!$D$5:$D$15)"
            elif lab == "Equity weight":
                v = f"=SUMPRODUCT({w},Inputs!$G$5:$G$15)"
            elif lab == "Illiquid weight":
                v = f"=SUMPRODUCT({w},Inputs!$H$5:$H$15)"
            elif lab == "Global financial crisis":
                v = f"=SUMPRODUCT({w},Stress!$B$11:$B$21)"
            elif lab == "COVID shock":
                v = f"=SUMPRODUCT({w},Stress!$C$11:$C$21)"
            elif lab.startswith("Inflation sell-off"):
                v = f"=SUMPRODUCT({w},Stress!$D$11:$D$21)"
            elif lab.startswith("US$150 oil (central"):
                v = f"=SUMPRODUCT({w},Stress!$E$11:$E$21)"
            elif lab.startswith("US$150 oil (severe"):
                v = f"=SUMPRODUCT({w},Oil!$C${oil_sev_row}:$C${oil_sev_row + 10})"
            elif lab == "One-way turnover":
                v = f"=SUMPRODUCT(ABS({w}-{CUR}))/2"
            elif lab.startswith("Transition cost"):
                v = f"=SUMPRODUCT(ABS({w}-{CUR}),Inputs!$B$54:$B$64)*Inputs!$B$22"
            elif lab.startswith("Compound advantage"):
                v = (f"={c}{rr['Expected return (compound)']}-(Risk!$N$56-Risk!$M$56^2/2)"
                     f"-{c}{rr['Transition cost (A$m)']}/Inputs!$B$22/Inputs!$B$69")
            else:
                te, tb, eq, il, ca = (rr["Tracking error (parametric)"], rr["Tracking error (bootstrap)"],
                                      rr["Equity weight"], rr["Illiquid weight"], W1)
                v = (f'=IF(AND({c}{te}<=Inputs!$B$27+0.00001,{c}{tb}<=Inputs!$B$27,{c}{eq}>=Inputs!$B$31-0.00001,'
                     f'{c}{eq}<=Inputs!$B$32+0.0005,{c}{il}<=Inputs!$B$33,{c}{ca}>=Inputs!$B$34-0.00001),"MET","NOT MET")')
            row.append(v)
        metrics.append(row)

    def mf(ri, ci):
        if ci == 0:
            return None
        lab = labels[ri]
        if lab.startswith("Sharpe"):
            return "0.00"
        if lab.startswith("Transition cost"):
            return "#,##0.0"
        if lab in ("Equity weight", "Illiquid weight", "One-way turnover") or lab.startswith(("Global", "COVID", "Inflation", "US$")):
            return PCT1
        return PCT2
    lastm = box(ws, top_m, 1, "9c. How the portfolios compare", ["Measure"] + NAMES, metrics, GREEN,
                fmts=mf, header_height=40)
    status_cf(ws, f"B{lastm}:F{lastm}")
    ws.row_dimensions[lastm].height = 27
    ws.cell(lastm, 1).alignment = Alignment(wrap_text=True, vertical="center")
    return {"rows": rr, "top_m": top_m, "total_row": W1 + 1, "status_row": lastm}


def summary(wb, ref):
    s = wb["Summary"]
    rr = ref["rows"]
    pick = [("Expected return (compound)", "Expected return (compound)", PCT2),
            ("Volatility", "Volatility", PCT2),
            ("Tracking error (bootstrap)", "Tracking error (bootstrap)", PCT2),
            ("5-year loss probability (bootstrap)", "5-year loss probability (bootstrap)", PCT1),
            ("Management fee (p.a.)", "Management fee (p.a.)", PCT2),
            ("Global financial crisis", "Global financial crisis", PCT1),
            ("US$150 oil (central case)", "US$150 oil (central case)", PCT1),
            ("Transition cost (A$m)", "Transition cost (A$m)", "#,##0.0"),
            ("Advantage vs current after costs", "Compound advantage vs current, after amortised costs", PCT2),
            ("Brief constraints", "Brief constraints (TE within 1.5% parametric and bootstrap, equity band, illiquid, cash)", None)]
    rows = []
    fm = []
    for lab, key, nf in pick:
        rows.append([lab] + [f"=Portfolio!{c}{rr[key]}" for c in "BCDEF"])
        fm.append(nf)
    last = box(s, 85, 1, "12. Asset-class scope: portfolios with and without some asset classes "
               "(Portfolio section 9)", ["Measure", "Recommended", "Alternatives optimised", "No alternatives",
                                         "Alternatives at benchmark", "Core five only"], rows, TEAL,
               fmts=lambda ri, ci: fm[ri] if ci else None, header_height=40)
    status_cf(s, f"B{last}:F{last}")


def run(wb):
    oil = wb["Oil"]
    sev = None
    for r in range(1, 400):
        v = oil.cell(r, 1).value
        if isinstance(v, str) and v.startswith("10b."):
            sev = r + 2
            break
    bootstrap(wb)
    ref = portfolio(wb, sev)
    summary(wb, ref)
    return ref
