"""Inputs, Portfolio, Validation and Checks additions: implementation, franking, integrity."""
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as L

from style import BLUE, GOLD, GREEN, GREY, NAVY, OLIVE, ORANGE, PURPLE, RED, TEAL, box, status_cf

ASSETS = ["Australian Equities", "World Equities", "Emerging Markets", "Listed Property",
          "Commodities", "Direct Property", "Hedge Funds", "Private Equity",
          "Australian Fixed Income", "World Fixed Income", "Cash"]
PCT1, PCT2 = "0.0%", "0.00%"
COV = "Risk!$B$41:$L$51"
BENCH = "Portfolio!$D$6:$D$16"
EQ = "Inputs!$G$5:$G$15"
ILLQ = "Inputs!$H$5:$H$15"


def te_col(rng):
    """Parametric tracking error of a column-form weight range against the benchmark."""
    return f"SQRT(SUMPRODUCT(MMULT({COV},{rng}-{BENCH}),{rng}-{BENCH}))"


def vol_col(rng):
    return f"SQRT(SUMPRODUCT(MMULT({COV},{rng}),{rng}))"


def ret_col(rng):
    return f"SUMPRODUCT(MMULT(Risk!$B$54:$L$54,{rng}))"


# ---------------------------------------------------------------------------------------------
def inputs(wb):
    ws = wb["Inputs"]
    costs = [0.001, 0.001, 0.0025, 0.002, 0.001, 0.015, 0.005, 0.01, 0.001, 0.0015, 0]
    notes = ["Brokerage + spread", "Brokerage + spread", "Wider spreads", "Spread",
             "Futures roll + spread", "Unlisted fund buy spread", "Fund buy spread",
             "Secondary or commitment costs", "Spread", "Spread (hedged)", "None"]
    rows = [[a, costs[k], notes[k]] for k, a in enumerate(ASSETS)]
    box(ws, 52, 1, "4. Implementation: one-way transaction costs (judgement, indicative)",
        ["Asset class", "Cost (one-way)", "Basis"], rows, ORANGE, fmts=[None, PCT2, None],
        inputs={(None, 1)})
    p = [
        ["Years to amortise transition costs", 5, "SIP horizon"],
        ["Direct property funded by end of year 1", 0.5, "Judgement: unlisted fund queues"],
        ["Direct property funded by end of year 2", 1.0, "Judgement"],
        ["Private equity funded by end of year 1", 1 / 3, "Commitments drawn over 3 years"],
        ["Private equity funded by end of year 2", 2 / 3, "Commitments drawn over 3 years"],
        ["Private equity funded by end of year 3", 1.0, "Commitments drawn over 3 years"],
        ["Unfunded direct property: share held in listed property", "=Forecast!L44", "Forecast factor model"],
        ["Unfunded private equity held in", "Cash", "Liquid reserve for capital calls"],
        ["Australian equity dividend yield (trailing)", "=Forecast!B41", "Forecast tab"],
        ["Share of dividends franked", 0.7, "Judgement: ASX franking level c.70%"],
        ["Company tax rate", 0.3, "Australian company tax rate"],
        ["Share of franking value the fund captures", 1.0, "Super fund taxed at 15% uses credits in full"],
    ]

    def pf(ri, ci):
        if ci != 1:
            return None
        return {0: "0", 7: None}.get(ri, PCT1)
    box(ws, 67, 1, "5. Transition and tax parameters", ["Parameter", "Value", "Source"], p, NAVY,
        fmts=pf, inputs={(ri, 1) for ri in range(len(p)) if ri not in (6, 7, 8)})
    # named rows used elsewhere
    return {"cost": "Inputs!$B$54:$B$64", "amort": "Inputs!$B$69", "dp1": "Inputs!$B$70",
            "dp2": "Inputs!$B$71", "pe1": "Inputs!$B$72", "pe2": "Inputs!$B$73", "pe3": "Inputs!$B$74",
            "dy": "Inputs!$B$77", "franked": "Inputs!$B$78", "tax": "Inputs!$B$79", "capture": "Inputs!$B$80"}


# ---------------------------------------------------------------------------------------------
def portfolio(wb, I):
    ws = wb["Portfolio"]
    # ---- 5. Transition costs ----
    top = 62
    rows = []
    for k, a in enumerate(ASSETS):
        r = top + 2 + k
        rows.append([a, f"=J{6 + k}", f"=Inputs!B{54 + k}", f"=ABS(B{r})*C{r}"])
    f, l = top + 2, top + 12
    rows += [["Total", f"=SUMPRODUCT(ABS(B{f}:B{l}))/2", None, f"=SUM(D{f}:D{l})"],
             ["Cost as % of fund", None, None, f"=D{l + 1}/Inputs!$B$22"],
             ["Amortised cost per year", None, None, f"=D{l + 2}/{I['amort']}"],
             ["Expected return advantage vs current (p.a.)", None, None, "=Risk!$N$58-Risk!$N$56"],
             ["Advantage after amortised transition costs", None, None, f"=D{l + 4}-D{l + 3}"],
             ["Status", None, None, f'=IF(D{l + 5}>=0,"PASS","REVIEW")']]

    def f5(ri, ci):
        if ci == 1:
            return "#,##0.0"
        if ci == 2:
            return PCT2
        if ci == 3:
            return "#,##0.00" if ri <= 11 else PCT2
        return None
    last = box(ws, top, 1, "5. Transition costs (A$m; one-way costs from Inputs section 4)",
               ["Asset class", "Trade (A$m)", "Cost rate", "Cost (A$m)"], rows, ORANGE, fmts=f5,
               bold_rows=(11,))
    ws[f"B{l + 1}"].number_format = "#,##0.0"
    ws[f"D{l + 1}"].number_format = "#,##0.00"
    status_cf(ws, f"D{last}")
    tc = {"total": f"Portfolio!$D${l + 1}", "pct": f"Portfolio!$D${l + 2}",
          "amort": f"Portfolio!$D${l + 3}", "after": f"Portfolio!$D${l + 5}", "status": f"Portfolio!$D${last}"}

    # ---- 6. Transition plan ----
    top6 = last + 3
    ph = ["Current", "Phase 1: liquid rebalance", "End of year 1", "End of year 2", "End of year 3 (target)"]
    dp_ramp = ["", 0, I["dp1"], I["dp2"], 1]
    pe_ramp = ["", 0, I["pe1"], I["pe2"], I["pe3"]]
    rows = []
    w0 = top6 + 2
    for k, a in enumerate(ASSETS):
        r = w0 + k
        row = [a, f"=C{6 + k}"]
        for j in range(1, 5):
            dpr, per = dp_ramp[j], pe_ramp[j]
            if a == "Direct Property":
                row.append(f"=$E{6 + k}*{dpr}")
            elif a == "Private Equity":
                row.append(f"=$E{6 + k}*{per}")
            elif a == "Listed Property":
                row.append(f"=$E{6 + k}+$E$11*(1-{dpr})*Inputs!$B$75")
            elif a == "Australian Fixed Income":
                row.append(f"=$E{6 + k}+$E$11*(1-{dpr})*(1-Inputs!$B$75)")
            elif a == "Cash":
                row.append(f"=$E{6 + k}+$E$13*(1-{per})")
            else:
                row.append(f"=$E{6 + k}")
        rows.append(row)
    wl = w0 + 10
    rows.append(["Total"] + [f"=SUM({c}{w0}:{c}{wl})" for c in "BCDEF"])
    metrics = [("Expected return (net)", ret_col), ("Volatility", vol_col), ("Tracking error", te_col)]
    for name, fn in metrics:
        rows.append([name] + [f"={fn(f'{c}${w0}:{c}${wl}')}" for c in "BCDEF"])
    rows.append(["Equity weight"] + [f"=SUMPRODUCT({c}{w0}:{c}{wl},{EQ})" for c in "BCDEF"])
    rows.append(["Illiquid weight"] + [f"=SUMPRODUCT({c}{w0}:{c}{wl},{ILLQ})" for c in "BCDEF"])
    rr = w0 + 11
    rows.append(["Brief constraints (TE 1.1% calibrated, equity band, illiquid, cash)"] +
                ["n/a: starting point"] + [f'=IF(AND({c}{rr + 3}<=Inputs!$B$45+0.00001,{c}{rr + 4}>=Inputs!$B$31-0.00001,'
                 f'{c}{rr + 4}<=Inputs!$B$32+0.00001,{c}{rr + 5}<=Inputs!$B$33,{c}{wl}>=Inputs!$B$34-0.00001),'
                 f'"MET","NOT MET")' for c in "CDEF"])

    def f6(ri, ci):
        if ci == 0:
            return None
        return PCT2 if ri in (12, 13, 14) else PCT1
    last6 = box(ws, top6, 1, "6. Transition plan: liquid assets move now; direct property and private "
                "equity are built over 1 to 3 years (undrawn amounts parked in liquid proxies)", ["Asset class"] + ph, rows, TEAL, fmts=f6,
                header_height=40, bold_rows=(11,))
    status_cf(ws, f"B{last6}:F{last6}")
    plan = {"status_row": last6, "first": w0}

    # ---- 7. Franking credits ----
    top7 = last6 + 3
    fr = f"{I['dy']}*{I['franked']}*{I['tax']}/(1-{I['tax']})*{I['capture']}"
    rows = [["Franking credit yield on Australian equities", None, None, f"={fr}", None]]
    r0 = top7 + 3
    names = [("Current fund", "C"), ("Benchmark", "D"), ("Recommended", "E")]
    for i, (n, c) in enumerate(names):
        r = r0 + i
        rows.append([n, f"={c}6", f"=Risk!$N${56 + i}", f"=B{r}*$D${top7 + 2}", f"=C{r}+D{r}"])
    rows.append(["Recommended less current", f"=B{r0 + 2}-B{r0}", f"=C{r0 + 2}-C{r0}",
                 f"=D{r0 + 2}-D{r0}", f"=E{r0 + 2}-E{r0}"])
    rows.append(["Status (advantage including franking)", None, None, None,
                 f'=IF(E{r0 + 3}>=0,"PASS","REVIEW")'])
    last7 = box(ws, top7, 1, "7. After-tax view: franking credits on Australian equities (the brief is "
                "pre-tax; this is a sensitivity)", ["Portfolio", "Australian equity weight",
                                                    "Expected return (pre-tax)", "Franking uplift",
                                                    "Return incl. franking"],
                rows, GOLD, fmts=[None, PCT1, PCT2, PCT2, PCT2], header_height=40)
    status_cf(ws, f"E{last7}")
    frank = {"yield": f"Portfolio!$D${top7 + 2}", "diff": f"Portfolio!$E${r0 + 3}",
             "uplift_diff": f"Portfolio!$D${r0 + 3}", "status": f"Portfolio!$E${last7}"}

    # ---- 8. Equity definition ----
    top8 = last7 + 3
    rows = []
    for i, (n, c) in enumerate(names):
        rows.append([n, f"=SUMPRODUCT({c}6:{c}16,{EQ})", f"={c}9", f"=B{top8 + 2 + i}+C{top8 + 2 + i}"])
    rows.append(["Band under the wider definition", None, None, f"=D{top8 + 3}-Inputs!$B$29"])
    rows.append([None, None, None, f"=D{top8 + 3}+Inputs!$B$29"])
    rows.append(["Recommended inside the wider band?", None, None,
                 f'=IF(AND(D{top8 + 4}>=D{top8 + 5}-0.00001,D{top8 + 4}<=D{top8 + 6}+0.00001),"PASS","REVIEW")'])
    last8 = box(ws, top8, 1, "8. Equity definition sensitivity: listed property counted as equity",
                ["Portfolio", "Equity (brief definition)", "Listed property", "Equity incl. listed property"],
                rows, PURPLE, fmts=[None, PCT1, PCT1, PCT1], header_height=28)
    ws[f"A{top8 + 5}"] = "Band under the wider definition: minimum"
    ws[f"A{top8 + 6}"] = "Band under the wider definition: maximum"
    status_cf(ws, f"D{last8}")
    eqd = {"rec": f"Portfolio!$D${top8 + 4}", "max": f"Portfolio!$D${top8 + 6}", "status": f"Portfolio!$D${last8}"}
    return tc, plan, frank, eqd


# ---------------------------------------------------------------------------------------------
def validation(wb):
    ws = wb["Validation"]
    sets = []
    m = wb["Methods"]
    for j, c in enumerate("DEFGH"):
        sets.append((f"Methods: {m[f'{c}5'].value}", f"Methods!{c}6:{c}16", True, "Inputs!$B$45"))
    for c, te in zip("BCDE", [None, "Inputs!$B$45", 0.0075, "Inputs!$B$27"]):
        sets.append((f"Sensitivity: {ws[f'{c}29'].value}", f"Validation!{c}30:{c}40", True, te))
    for c in "BC":
        sets.append((f"Out-of-sample: {ws[f'{c}16'].value}", f"Validation!{c}17:{c}27", True, None))
    p = wb["Portfolio"]
    for c, te in zip("CDEFG", ["Inputs!$B$27", None, "Inputs!$B$45", "Inputs!$B$45", None]):
        sets.append((f"Constraint cost: {p[f'{c}42'].value}", f"Portfolio!{c}43:{c}53", c not in "EG", te))
    rows = []
    top = 85
    for name, rng, eqband, te in sets:
        r = top + 2 + len(rows)
        col = rng.split("!")[1]
        a, b = col.split(":")
        rows.append([name, f"=SUM({rng})", f"=MIN({rng})", f"={rng.split('!')[0]}!{b}",
                     f"=SUMPRODUCT({rng},{EQ})", f"={te_col(rng)}", (f"={te}" if isinstance(te, str) else te) if te is not None else "n/a",
                     f'=IF(AND(ABS(B{r}-1)<0.0005,C{r}>=-0.000001,D{r}>=Inputs!$B$34-0.00001,'
                     + (f'E{r}>=Inputs!$B$31-0.00001,E{r}<=Inputs!$B$32+0.0005,' if eqband else '')
                     + (f'F{r}<=G{r}+0.0002' if te is not None else 'TRUE') + '),"PASS","REVIEW")'])
    # frontier sets (ten points each): worst case across the points
    for title, r0, te in (("Frontier: no TE limit (10 points)", 113, None),
                          ("Frontier: with TE limit (10 points)", 131, "Inputs!$B$45")):
        cols = [L(2 + i) for i in range(10)]
        r = top + 2 + len(rows)
        sums = ",".join(f"ABS(SUM({c}{r0}:{c}{r0 + 10})-1)" for c in cols)
        cash = ",".join(f"{c}{r0 + 10}" for c in cols)
        eqs = [f"SUMPRODUCT({c}{r0}:{c}{r0 + 10},{EQ})" for c in cols]
        tes = ",".join(te_col(f"{c}{r0}:{c}{r0 + 10}") for c in cols)
        rows.append([title, f"=1+MAX({sums})", f"=MIN(B{r0}:K{r0 + 10})", f"=MIN({cash})",
                     f"=MAX({','.join(eqs)})", f"=MAX({tes})", f"={te}" if te else "n/a",
                     f'=IF(AND(ABS(B{r}-1)<0.0005,C{r}>=-0.000001,D{r}>=Inputs!$B$34-0.00001,'
                     f'E{r}<=Inputs!$B$32+0.0005,' + (f'F{r}<=G{r}+0.0002' if te else 'TRUE') +
                     '),"PASS","REVIEW")'])

    def fv(ri, ci):
        return {1: "0.0000", 2: PCT2, 3: PCT1, 4: PCT1, 5: PCT2, 6: PCT2}.get(ci)
    last = box(ws, top, 1, "7. Integrity of pasted Solver outputs (re-tested against current inputs)",
               ["Weight set", "Sum of weights", "Smallest weight", "Cash", "Equity (worst point)",
                "Tracking error (worst point)", "TE limit", "Status"], rows, PURPLE, fmts=fv,
               header_height=40)
    status_cf(ws, f"H{top + 2}:H{last}")
    ws[f"A{last + 1}"] = ("Solver outputs are pasted values. If any input changes, re-run Solver for the "
                          "rows marked REVIEW and paste the new weights.")
    ws[f"A{last + 1}"].font = ws[f"A{last + 1}"].font.copy(name="Arial", size=9, italic=True, color="FF595959")
    assert last + 1 < 111
    return {"range": f"Validation!$H${top + 2}:$H${last}"}


# ---------------------------------------------------------------------------------------------
def checks(wb, C, perf, oilref, tc, plan, frank, eqd, integ):
    ws = wb["Checks"]
    existing = []
    for r in range(7, 31):
        existing.append(([ws.cell(r, c).value for c in range(1, 5)],
                         [ws.cell(r, c).number_format for c in range(1, 5)]))
    perf_rows, attrib, blocks = perf
    act = C["Active (net)"]
    new = [
        (["Performance panel complete (monthly, Apr-2006 to Mar-2026)", f"=COUNT(Performance!{act}6:{act}245)", 240,
          '=IF(B{r}=C{r},"PASS","FAIL")'], "0"),
        (["Attribution reconciles to active return (5 years)",
          f"=ABS(Performance!F{attrib['total']}-Performance!F{attrib['arith']})", 0,
          '=IF(B{r}<0.0000001,"PASS","FAIL")'], "0.0000"),
        (["Attribution reconciles to active return (20 years)",
          f"=ABS(Performance!H{attrib['total']}-Performance!H{attrib['arith']})", 0,
          '=IF(B{r}<0.0000001,"PASS","FAIL")'], "0.0000"),
        (["Max weight gap formulas evaluate (Risk P59:P62)", "=COUNT(Risk!P59:P62)", 4,
          '=IF(B{r}=C{r},"PASS","FAIL")'], "0"),
        (["AU 3-month rate gaps interpolated (zeros in Oil panel)", "=COUNTIF(Oil!S5:S701,0)", 0,
          '=IF(B{r}=C{r},"PASS","FAIL")'], "0"),
        (["Oil robustness regressions have 30+ observations",
          f"=MIN(Oil!H{oilref['rob_first']}:H{oilref['rob_first'] + 11})", 30, '=IF(B{r}>=C{r},"PASS","FAIL")'], "0"),
        (["Asset-class oil regressions have 30+ observations",
          f"=MIN(Oil!G{oilref['asset_first']}:G{oilref['asset_first'] + 10})", 30, '=IF(B{r}>=C{r},"PASS","FAIL")'], "0"),
        (["2008 oil episode prices found", f"=COUNT(Oil!D{oilref['e0']}:E{oilref['e0']})", 2,
          '=IF(B{r}=C{r},"PASS","FAIL")'], "0"),
        (["Pasted Solver outputs still meet constraints", f'=COUNTIF({integ["range"]},"REVIEW")', 0,
          '=IF(B{r}=C{r},"PASS","REVIEW")'], "0"),
        (["Transition plan phases meet constraints",
          f'=COUNTIF(Portfolio!C{plan["status_row"]}:F{plan["status_row"]},"NOT MET")', 0,
          '=IF(B{r}=C{r},"PASS","REVIEW")'], "0"),
        (["Robust vs resampled largest weight gap within limit", "=Methods!D33", "=Methods!E33",
          '=IF(B{r}<=C{r},"PASS","REVIEW")'], "0.0%"),
        (["Return advantage survives transition costs", f"={tc['after']}", 0,
          '=IF(B{r}>=C{r},"PASS","REVIEW")'], "0.00%"),
        (["Return advantage survives franking adjustment", f"={frank['diff']}", 0,
          '=IF(B{r}>=C{r},"PASS","REVIEW")'], "0.00%"),
        (["Equity inside band if listed property counts as equity", f"={eqd['rec']}", f"={eqd['max']}",
          '=IF(B{r}<=C{r}+0.00001,"PASS","REVIEW")'], "0.0%"),
    ]
    rows, fmts = [], {}
    for i, (vals, nf) in enumerate(existing):
        rows.append(vals)
        fmts[i] = nf
    for j, (vals, nf) in enumerate(new):
        r = 7 + len(rows)
        rows.append([v.replace("{r}", str(r)) if isinstance(v, str) else v for v in vals])
        fmts[len(rows) - 1] = [None, nf, nf, None]
    last = box(ws, 5, 1, "Tests", ["Test", "Value", "Limit", "Status"], rows, GREEN,
               fmts=lambda ri, ci: fmts[ri][ci])
    for r in range(7, last + 1):
        ws[f"D{r}"].alignment = Alignment(horizontal="center")
    ws["B4"] = (f'=IF(COUNTIF(D7:D{last},"FAIL")=0,"ALL CHECKS PASS","FAIL: "&COUNTIF(D7:D{last},"FAIL")'
                f'&" check(s)")')
    # existing CF covers D7:D30; extend to the new rows
    status_cf(ws, f"D31:D{last}")
    ws[f"A{last + 2}"] = ("REVIEW marks a judgement to disclose in the report, not a model error. "
                          "The master flag counts FAIL only.")
    ws[f"A{last + 2}"].font = ws[f"A{last + 2}"].font.copy(name="Arial", size=9, italic=True, color="FF595959")
    ws[f"A{last + 3}"] = f'=COUNTIF(D7:D{last},"REVIEW")&" item(s) flagged REVIEW"'
    ws[f"A{last + 3}"].font = ws[f"A{last + 3}"].font.copy(name="Arial", size=9, bold=True, color="FF806000")
    return {"review": f"Checks!$A${last + 3}", "last": last}
