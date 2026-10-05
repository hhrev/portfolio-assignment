"""Summary boxes and Report figures data blocks for the v3.4 additions."""
from copy import copy

from openpyxl.styles import Border, Side

from style import GREY, NAVY, OLIVE, ORANGE, BROWN, box, status_cf

PCT1, PCT2 = "0.0%", "0.00%"
ASSETS = ["Australian Equities", "World Equities", "Emerging Markets", "Listed Property",
          "Commodities", "Direct Property", "Hedge Funds", "Private Equity",
          "Australian Fixed Income", "World Fixed Income", "Cash"]


def summary(wb, C, perf, oilref, tc, plan, frank, eqd, chk):
    ws = wb["Summary"]
    pr, attrib, blocks = perf
    P = "Performance!"
    rows = [
        ["Return, Mar-21 to Mar-26", f"={P}B{pr['Annualised return, net of fees']}",
         f"={P}C{pr['Annualised return, net of fees']}", f"={P}D{pr['Annualised return, net of fees']}"],
        ["Return, 20 years", f"={P}E{pr['Annualised return, net of fees']}",
         f"={P}F{pr['Annualised return, net of fees']}", f"={P}G{pr['Annualised return, net of fees']}"],
        ["Sharpe ratio, 5 years", f"={P}B{pr['Sharpe ratio (excess over bank bills)']}",
         f"={P}C{pr['Sharpe ratio (excess over bank bills)']}", f"={P}D{pr['Sharpe ratio (excess over bank bills)']}"],
        ["Tracking error, 5 years", f"={P}B{pr['Tracking error']}", None, None],
        ["Information ratio, 5 years", f"={P}B{pr['Information ratio']}", None, None],
        ["Maximum drawdown, 20 years", f"={P}E{pr['Maximum drawdown']}", f"={P}F{pr['Maximum drawdown']}",
         f"={P}G{pr['Maximum drawdown']}"],
        ["Down-market capture, 20 years", f"={P}E{pr['Down-market capture']}", 1, None],
        ["Active return, 2006-11 block", f"={P}D{blocks['blocks_first']}", None, None],
        ["Active return, 2021-26 block", f"={P}D{blocks['blocks_last']}", None, None],
    ]

    def f9(ri, ci):
        if ci == 0:
            return None
        return "0.00" if ri in (2, 4, 6) else PCT2
    last9 = box(ws, 56, 8, "9. Past performance (monthly rebalancing, net of fees)",
                ["Measure", "Current", "Benchmark", "Difference"], rows, OLIVE, fmts=f9)

    o = oilref
    rb = o["rob_first"]
    imp = o["impact"]["rec"]
    of = o["opt_first"]
    rows = [
        ["Australian equities: 12m oil beta since 1986", f"=Oil!E{rb + 4}"],
        ["AU 5-year yield: 12m oil beta since 1986 (pp)", f"=Oil!E{rb + 10}"],
        ["Recommended: regression impact, 12 months", f"=Oil!B{imp}"],
        ["Recommended: episode central case", f"=Oil!F{imp}"],
        ["Recommended: severe case (2008 scaled)", f"=Oil!E{imp}"],
        ["Mitigation: commodities +2% from Aus equities, loss reduced", f"=Oil!H{of + 3}"],
        ["Mitigation: same option, tracking error", f"=Oil!D{of + 3}"],
    ]
    rows = [[r[0], None, None, r[1]] for r in rows]
    last10 = box(ws, last9 + 3, 8, "10. Oil: robustness and mitigation", ["Measure", None, None, "Value"],
                 rows, BROWN, fmts=lambda ri, ci: ("0.000" if ri == 0 else "0.00" if ri == 1 else PCT2) if ci == 3 else None)

    rows = [
        ["Transition cost (A$m, one-way)", f"={tc['total']}", None],
        ["Transition cost (% of fund)", f"={tc['pct']}", None],
        ["Amortised cost per year", f"={tc['amort']}", None],
        ["Expected return advantage after costs", f"={tc['after']}", f"={tc['status']}"],
        ["Franking uplift: recommended less current", f"={frank['uplift_diff']}", None],
        ["Return advantage including franking", f"={frank['diff']}", f"={frank['status']}"],
        ["Equity incl. listed property (recommended)", f"={eqd['rec']}", f"={eqd['status']}"],
        ["Transition phases meeting constraints",
         f'=COUNTIF(Portfolio!C{plan["status_row"]}:F{plan["status_row"]},"MET")&" of 4"', None],
        ["Model checks flagged for disclosure", f'=COUNTIF(Checks!D7:D{chk["last"]},"REVIEW")', None],
    ]

    def f11(ri, ci):
        if ci != 1:
            return None
        return {0: "#,##0.0", 7: None, 8: "0"}.get(ri, PCT2)
    last11 = box(ws, 72, 1, "11. Implementation and after-tax view", ["Measure", "Value", "Status"],
                 rows, ORANGE, fmts=f11)
    status_cf(ws, f"C74:C{last11}")


def figures(wb, C, perf):
    ws = wb["Report figures"]
    pr, attrib, blocks = perf
    # index rows 18 and 19
    for r in (18, 19):
        for c in range(1, 7):
            src = ws.cell(17, c)
            dst = ws.cell(r, c)
            dst._style = copy(src._style)
    ws["A18"], ws["B18"], ws["F18"] = "Figure 14", "Rolling 5-year active return (net of fees)", "Performance"
    ws["A19"], ws["B19"], ws["F19"] = "Figure 15", "Attribution of active return", "Performance"
    thin = Side("thin", color="FFD0D0D0")
    for r in (17, 18):
        for c in range(1, 7):
            b = ws.cell(r, c).border
            ws.cell(r, c).border = Border(left=b.left, right=b.right, top=b.top, bottom=thin)

    def title_row(r, text):
        for c in range(1, 13):
            ws.cell(r, c)._style = copy(ws.cell(332, c)._style)
        ws.cell(r, 1).value = text

    # Figure 14: rolling 60-month active return
    title_row(358, "Figure 14. Rolling 5-year active return, net of fees")
    ws["AB358"] = "Chart data, figure 14"
    ws["AB358"]._style = copy(ws["AB254"]._style)
    for c, h in zip(("AB", "AC", "AD"), ("Window end", "Active return (current less benchmark)", "Zero")):
        ws[f"{c}359"] = h
        ws[f"{c}359"]._style = copy(ws["AC255"]._style if c != "AB" else ws["AB255"]._style)
    ra, me = C["Rolling 60m active"], C["Month end"]
    n = 0
    for i, pr_row in enumerate(range(65, 246)):
        r = 360 + i
        ws[f"AB{r}"] = f"=Performance!{me}{pr_row}"
        ws[f"AC{r}"] = f"=Performance!{ra}{pr_row}"
        ws[f"AD{r}"] = 0
        ws[f"AB{r}"]._style = copy(ws["AB256"]._style)
        ws[f"AC{r}"]._style = copy(ws["AC256"]._style)
        ws[f"AD{r}"]._style = copy(ws["AC256"]._style)
        n += 1
    fig14 = {"cat": f"$AB$360:$AB${359 + n}", "s1": "$AC$359", "v1": f"$AC$360:$AC${359 + n}",
             "s2": "$AD$359", "v2": f"$AD$360:$AD${359 + n}", "n": n}

    # Figure 15: attribution
    title_row(384, "Figure 15. Attribution of active return, net of fees")
    ws["N384"] = "Chart data, figure 15"
    ws["N384"]._style = copy(ws["N332"]._style)
    for c, h in zip(("N", "O", "P"), ("Source", "Mar-21 to Mar-26", "20 years")):
        ws[f"{c}385"] = h
        ws[f"{c}385"]._style = copy(ws[f"{c}333"]._style)
    labels = ASSETS + ["Fee difference"]
    for i, lab in enumerate(labels):
        r = 386 + i
        src = attrib["first"] + i if i < 11 else attrib["fee"]
        ws[f"N{r}"] = lab
        ws[f"O{r}"] = f"=Performance!F{src}"
        ws[f"P{r}"] = f"=Performance!H{src}"
        for c in "NOP":
            ws[f"{c}{r}"]._style = copy(ws[f"{c}334"]._style)
    fig15 = {"cat": "$N$386:$N$397", "s1": "$O$385", "v1": "$O$386:$O$397", "s2": "$P$385",
             "v2": "$P$386:$P$397", "n": 12}
    return fig14, fig15
