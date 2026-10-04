"""Oil tab upgrades (criterion 3)."""
from openpyxl.utils import get_column_letter as L

from style import BLUE, BROWN, GOLD, GREY, OLIVE, PURPLE, RED, TEAL, box, status_cf

ASSETS = ["Australian Equities", "World Equities", "Emerging Markets", "Listed Property",
          "Commodities", "Direct Property", "Hedge Funds", "Private Equity",
          "Australian Fixed Income", "World Fixed Income", "Cash"]
PCT1, PCT2 = "0.0%", "0.00%"
R0, R1 = 5, 701          # panel rows (Jan-1968 to Jan-2026)
R86 = 221                # Jan-1986: start of the post-1986 sample
X0 = 33                  # column AG: start of the panel extension


def ext_cols():
    names = ["Month end", "AU 5-year yield (%)", "1m yield change (pp)", "12m yield change (pp)",
             "60m yield change (pp)"]
    names += [f"ln {a}" for a in ASSETS]
    for h in (1, 12, 60):
        names += [f"{h}m {a} (log)" for a in ASSETS]
    return {n: L(X0 + i) for i, n in enumerate(names)}, names


def run(wb):
    ws = wb["Oil"]

    # ---- data fix: the FRED 3-month rate has a zero for Nov-1969 (missing value) ----
    src = "'Additional data'!$E$2:$E$699"
    for r in range(R0, R1 + 1):
        m = f"MATCH(DATE(YEAR(N{r}),MONTH(N{r}),1),'Additional data'!$D$2:$D$699,0)"
        ws[f"S{r}"] = (f"=IF(INDEX({src},{m})=0,(INDEX({src},{m}-1)+INDEX({src},{m}+1))/2,"
                       f"INDEX({src},{m}))")

    # ---- panel extension ----
    X, names = ext_cols()
    rows = []
    for r in range(R0, R1 + 1):
        y = X["AU 5-year yield (%)"]
        row = [f"=N{r}",
               f"=IFERROR(INDEX('RBA 5Y yield'!$B$5:$B$661,MATCH(N{r},'RBA 5Y yield'!$A$5:$A$661,1))*100,\"\")"]
        for h in (1, 12, 60):
            row.append(f"=IF(AND(ISNUMBER({y}{r}),ISNUMBER({y}{r - h})),{y}{r}-{y}{r - h},\"\")"
                       if r - h >= R0 else "")
        for k, a in enumerate(ASSETS):
            dc = L(2 + k)
            row.append(f"=IFERROR(LN(INDEX(Data!${dc}$2:${dc}$497,MATCH(N{r},Data!$A$2:$A$497,1))),\"\")")
        for h in (1, 12, 60):
            for a in ASSETS:
                lv = X[f"ln {a}"]
                row.append(f"=IF(AND(ISNUMBER({lv}{r}),ISNUMBER({lv}{r - h})),{lv}{r}-{lv}{r - h},\"\")"
                           if r - h >= R0 else "")
        rows.append(row)

    def xfmt(ri, ci):
        if ci == 0:
            return "mmm-yy"
        if ci <= 4:
            return "0.00"
        if ci <= 15:
            return "0.0000"
        return "0.000"
    box(ws, 3, X0, "Panel extension: AU 5-year yield (RBA 5Y yield tab) and asset-class log levels "
        "and changes (Data tab)", names, rows, GREY, fmts=xfmt, header_height=54)
    for i in range(len(names)):
        ws.column_dimensions[L(X0 + i)].width = 10 if i == 0 else 11
    ws.column_dimensions["AF"].width = 3

    def rr(col, a=None):
        return f"${col}${a or R0}:${col}${R1}"

    # ---- 7. Robustness ----
    series = [("US equities", "U", "Y", "AC", True), ("Australian equities", "V", "Z", "AD", True),
              ("AU 3-month rate (pp)", "W", "AA", "AE", False),
              ("AU 5-year yield (pp)", X["1m yield change (pp)"], X["12m yield change (pp)"],
               X["60m yield change (pp)"], False)]
    oilc = {1: "T", 12: "X", 60: "AB"}
    top = 89
    rows = []
    for name, c1, c12, c60, is_eq in series:
        for h, cy in ((1, c1), (12, c12), (60, c60)):
            r = top + 2 + len(rows)
            ox = oilc[h]
            full_y, full_x = rr(cy), rr(ox)
            sub_y, sub_x = rr(cy, R86), rr(ox, R86)
            resp = f"=EXP(E{r}*$B$7)-1" if is_eq else f"=E{r}*$B$7"
            rows.append([name, h, f"=SLOPE({full_y},{full_x})",
                         f"=CORREL({full_y},{full_x})*SQRT((COUNT({full_y})-2)/(1-RSQ({full_y},{full_x})))/SQRT(B{r})",
                         f"=SLOPE({sub_y},{sub_x})",
                         f"=CORREL({sub_y},{sub_x})*SQRT((H{r}-2)/(1-G{r}))/SQRT(B{r})",
                         f"=RSQ({sub_y},{sub_x})", f"=COUNT({sub_y})", resp])

    def f7(ri, ci):
        return {1: "0", 2: "0.000", 3: "0.00", 4: "0.000", 5: "0.00", 6: "0.000", 7: "0",
                8: (PCT1 if ri < 6 else "0.00")}.get(ci)
    last7 = box(ws, top, 1, "7. Robustness: post-1986 sample (oil price deregulated) and the 5-year "
                "bond yield", ["Series", "Horizon (months)", "Beta: full sample", "t-stat (adj.): full",
                               "Beta: since 1986", "t-stat (adj.): since 1986", "R-squared: since 1986",
                               "Obs: since 1986", "Response to scenario (since 1986)"],
                rows, BROWN, fmts=f7, header_height=40)
    rob = {"first": top + 2}

    # ---- 8. Asset-class sensitivity (regression approach) ----
    top8 = last7 + 3
    rows = []
    for k, a in enumerate(ASSETS):
        r = top8 + 2 + k
        cy = {h: X[f"{h}m {a} (log)"] for h in (1, 12, 60)}
        rows.append([a] + [f"=SLOPE({rr(cy[h])},{rr(oilc[h])})" for h in (1, 12, 60)] +
                    [f"=CORREL({rr(cy[12])},{rr(oilc[12])})*SQRT((G{r}-2)/(1-F{r}))/SQRT(12)",
                     f"=RSQ({rr(cy[12])},{rr(oilc[12])})", f"=COUNT({rr(cy[12])})"] +
                    [f"=EXP({c}{r}*$B$7)-1" for c in "BCD"])
    f8 = top8 + 2

    def f8f(ri, ci):
        return {1: "0.000", 2: "0.000", 3: "0.000", 4: "0.00", 5: "0.000", 6: "0"}.get(ci, PCT1 if ci else None)
    last8 = box(ws, top8, 1, "8. Asset-class sensitivity to oil (log return on log oil change, all "
                "available months)", ["Asset class", "Beta: 1 month", "Beta: 12 months",
                                      "Beta: 60 months", "t-stat (adj.): 12m", "R-squared: 12m",
                                      "Obs: 12m", "Response: 1 month", "Response: 12 months",
                                      "Response: 60 months"], rows, PURPLE, fmts=f8f, header_height=40)

    # ---- 10. 2008 episode (built before section 9 so 9 can reference it) ----
    top9 = last8 + 3
    top10 = top9 + 8
    e0 = top10 + 2
    rows = [["Oil spike 2008", "=DATE(2007,12,31)", "=DATE(2008,6,30)",
             f"=INDEX('Additional data'!$B$2:$B$965,MATCH(DATE(YEAR(B{e0}),MONTH(B{e0}),1),'Additional data'!$A$2:$A$965,0))",
             f"=INDEX('Additional data'!$B$2:$B$965,MATCH(DATE(YEAR(C{e0}),MONTH(C{e0}),1),'Additional data'!$A$2:$A$965,0))",
             f"=E{e0}/D{e0}-1", f"=$B$7/LN(1+F{e0})"],
            ["Include 2008 in central case (1 = yes; default 0, used as a severe case)", 0, None, None, None, None, None]]

    def f10a(ri, ci):
        return {1: "mmm-yy", 2: "mmm-yy", 3: "0.00", 4: "0.00", 5: PCT1, 6: "0.00"}.get(ci) if ri == 0 else ("0" if ci == 1 else None)
    last10a = box(ws, top10, 1, "10. Third episode: 2008 oil spike (WTI to US$134 monthly average; "
                  "demand-led and overlaps the early GFC)", ["Episode", "Start", "End", "WTI start",
                                                             "WTI end", "WTI change", "Scale to scenario"],
                  rows, RED, fmts=f10a, inputs={(0, 1), (0, 2), (1, 1)})
    flag = f"$B${e0 + 1}"
    top10b = last10a + 2
    rows = []
    for k, a in enumerate(ASSETS):
        r = top10b + 2 + k
        dc = L(2 + k)
        rows.append([a, f"=INDEX(Data!{dc}$2:{dc}$497,MATCH($C${e0},Data!$A$2:$A$497,0))/"
                        f"INDEX(Data!{dc}$2:{dc}$497,MATCH($B${e0},Data!$A$2:$A$497,0))-1",
                     (0 if a == "Cash" else f"=B{r}*$G${e0}"), f"=AVERAGE(D{28 + k}:E{28 + k})",
                     f"=AVERAGE(D{28 + k}:E{28 + k},C{r})"])
    a10 = top10b + 2
    last10 = box(ws, top10b, 1, "10b. Asset returns: 2008 episode and the central case",
                 ["Asset class", "2008 as happened", "2008 scaled", "Central case: 1990 and 2022",
                  "Central case: 1990, 2022 and 2008"], rows, RED, fmts=[None, PCT1, PCT1, PCT1, PCT1],
                 header_height=40)
    # central case used everywhere (Stress, Summary, section 6) switches on the flag
    for k in range(11):
        ws[f"F{28 + k}"] = f"=IF({flag}=1,E{a10 + k},D{a10 + k})"
    ws["F27"] = "Central case"
    ws["G38"] = "Not scaled"
    ws["A22"] = "3. Episode study: oil price shocks, scaled to US$150 (2008 added in section 10)"

    # ---- 9. Portfolio impact by horizon ----
    prow = {"Current fund": "Portfolio!$C$6:$C$16", "Benchmark": "Portfolio!$D$6:$D$16",
            "Recommended": "Portfolio!$E$6:$E$16"}
    rows = []
    for name, w in prow.items():
        rows.append([name] + [f"=SUMPRODUCT({c}${f8}:{c}${f8 + 10},{w})" for c in "HIJ"] +
                    [f"=SUMPRODUCT(D$28:D$38,{w})", f"=SUMPRODUCT(E$28:E$38,{w})",
                     f"=SUMPRODUCT(C${a10}:C${a10 + 10},{w})", f"=SUMPRODUCT(F$28:F$38,{w})"])
    last9 = box(ws, top9, 1, "9. US$150 oil: portfolio impact by horizon, regression approach vs "
                "episode approach", ["Portfolio", "Regression: 1 month", "Regression: 12 months",
                                     "Regression: 60 months", "Episode: 1990 scaled",
                                     "Episode: 2022 scaled", "Episode: 2008 scaled", "Central case"],
                rows, BROWN, fmts=[None] + [PCT1] * 7, header_height=40)
    imp = {"cur": top9 + 2, "bench": top9 + 3, "rec": top9 + 4}
    assert last9 < top10

    # ---- 11. Mitigation options ----
    top11 = last10 + 3
    opts = [("Recommended (no change)", {}),
            ("Commodities +2%, world equities -2%", {"Commodities": 0.02, "World Equities": -0.02}),
            ("Commodities +1%, world equities -1%", {"Commodities": 0.01, "World Equities": -0.01}),
            ("Commodities +2%, Australian equities -2%", {"Commodities": 0.02, "Australian Equities": -0.02}),
            ("Commodities +4%, Australian equities -4%", {"Commodities": 0.04, "Australian Equities": -0.04}),
            ("Shorter duration: Australian fixed income -5%, cash +5%",
             {"Australian Fixed Income": -0.05, "Cash": 0.05}),
            ("Combined: commodities +2% (from world equities), AFI -3% to cash",
             {"Commodities": 0.02, "World Equities": -0.02, "Australian Fixed Income": -0.03, "Cash": 0.03})]
    wtop = top11 + 2 + len(opts) + 3   # row-form weights block below the results
    rows = []
    for i, (name, tilt) in enumerate(opts):
        wr = wtop + 2 + i
        ar = wtop + 2 + len(opts) + i
        w = f"$B${wr}:$L${wr}"
        a = f"$B${ar}:$L${ar}"
        r = top11 + 2 + i
        rows.append([name, f"=SUMPRODUCT(Risk!$B$54:$L$54,{w})",
                     f"=SQRT(SUMPRODUCT(MMULT({w},Risk!$B$41:$L$51),{w}))",
                     f"=SQRT(SUMPRODUCT(MMULT({a},Risk!$B$41:$L$51),{a}))",
                     f"=SUMPRODUCT(MMULT({w},$F$28:$F$38))",
                     f"=SUMPRODUCT(MMULT({w},$I${f8}:$I${f8 + 10}))",
                     f"=B{r}-B${top11 + 2}", f"=E{r}-E${top11 + 2}",
                     f'=IF(AND(D{r}<=Inputs!$B$45+0.00001,MIN({w})>=0,SUMPRODUCT(MMULT({w},Inputs!$G$5:$G$15))<=Inputs!$B$32,SUMPRODUCT(MMULT({w},Inputs!$G$5:$G$15))>=Inputs!$B$31),"MET","NOT MET")'])
    last11 = box(ws, top11, 1, "11. Mitigation options compared (parametric, same risk model as the "
                 "recommendation)", ["Option", "Expected return", "Volatility", "Tracking error",
                                     "Oil: episode central case", "Oil: 12-month regression",
                                     "Return cost vs recommended", "Oil loss reduced by",
                                     "TE (1.1% calibrated limit) and equity band"], rows, OLIVE,
                 fmts=[None, PCT2, PCT2, PCT2, PCT1, PCT1, PCT2, PCT1, None], header_height=40)
    status_cf(ws, f"I{top11 + 2}:I{last11}")
    rows = []
    for name, tilt in opts:
        rows.append([name] + [f"=Risk!{L(2 + k)}$58+{tilt.get(a, 0)}" for k, a in enumerate(ASSETS)])
    for i, (name, tilt) in enumerate(opts):
        wr = wtop + 2 + i
        rows.append([f"Active: {name}"] + [f"={L(2 + k)}{wr}-Risk!{L(2 + k)}$57" for k in range(11)])
    box(ws, wtop, 1, "11b. Option weights (row form) and active weights vs benchmark",
        ["Weights"] + ASSETS, rows, GREY, fmts=[None] + [PCT1] * 11, header_height=40)
    for r in range(wtop + 2, wtop + 2 + len(opts)):
        for k, a in enumerate(ASSETS):
            t = opts[r - wtop - 2][1].get(a, 0)
            if t:
                ws[f"{L(2 + k)}{r}"].font = ws[f"{L(2 + k)}{r}"].font.copy(color="FF0000FF")
    last11b = wtop + 1 + 2 * len(opts)

    # ---- 12. Other tools (qualitative) ----
    top12 = last11b + 3
    tools = [
        ["Oil call options or equity put overlay",
         "Pays off directly in a spike; caps the loss without changing the SAA",
         "Premium is a certain drag (typically 1-3% of notional a year); needs a derivatives mandate"],
        ["Inflation-linked bonds in place of some nominal bonds",
         "Oil shocks feed headline inflation; real-yield bonds protect purchasing power",
         "Not in the asset universe or fee schedule; lower expected return than nominal bonds today"],
        ["Energy and resources tilt within Australian equities",
         "ASX energy and miners have positive oil betas; keeps the equity weight unchanged",
         "Concentration and active risk; adds to tracking error"],
        ["Keep world equities unhedged",
         "The AUD tends to fall in global risk-off episodes, cushioning offshore assets in A$",
         "Currency volatility in normal times; the AUD is partly a commodity currency"],
        ["Shorter fixed income duration",
         "Rates rose with oil in 2022 (section 7, 12-month yield beta); less duration means smaller bond losses",
         "Gives up term premium and the hedge bonds provide in demand-led recessions"],
        ["Liquidity buffer and rebalancing rule",
         "Cash at the 2% floor plus liquid assets let the fund rebalance into a sell-off",
         "Small cash drag"],
    ]
    tools = [t + [None] * 6 for t in tools]
    lastq = box(ws, top12, 1, "12. Other mitigation tools (not modelled; discuss in the report)",
                ["Tool", "How it helps", None, None, None, "Cost or drawback"] + [None] * 3,
                [[t[0], t[1], None, None, None, t[2], None, None, None] for t in tools], GOLD)
    from openpyxl.styles import Alignment
    for r in range(top12 + 1, lastq + 1):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
        for c in ("A", "B", "F"):
            ws[f"{c}{r}"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        if r > top12 + 1:
            ws.row_dimensions[r].height = 40
    return {"rob_first": rob["first"], "asset_first": f8, "impact": imp, "flag": flag,
            "opt_first": top11 + 2, "a10": a10, "e0": e0}
